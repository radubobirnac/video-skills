#!/usr/bin/env python3
"""Local, read-only Amazon export validation. Python 3.10+, optional openpyxl."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, getcontext
from pathlib import Path
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

VERSION = "1.0.0"
getcontext().prec = 50
MAX_BYTES = 100 * 1024 * 1024
MAX_ROWS = 250_000
MARKETS = {"US": "USD", "UK": "GBP", "CA": "CAD"}
KINDS = {"auto", "campaign", "search_term", "targeting", "advertised_product", "business", "bulk"}
COUNTS = {"impressions", "clicks", "orders", "units", "sessions", "ordered_units", "order_items"}
MONEY = {"spend", "sales", "ordered_sales", "bid", "budget"}
METRICS = COUNTS | MONEY
IDS = {"campaign_id", "ad_group_id", "keyword_id", "target_id"}
ALIASES = {
    "campaign_id": ["Campaign ID", "campaignId"],
    "campaign_name": ["Campaign Name", "campaignName", "Campaign"],
    "ad_group_id": ["Ad Group ID", "adGroupId"],
    "ad_group_name": ["Ad Group Name", "adGroupName"],
    "keyword_id": ["Keyword ID", "keywordId"],
    "target_id": ["Target ID", "targetingId", "productTargetingId"],
    "search_term": ["Customer Search Term", "Search Term", "searchTerm"],
    "target": ["Targeting", "Keyword", "Keyword Text", "keywordText", "Targeting Expression"],
    "match_type": ["Match Type", "matchType"],
    "advertised_asin": ["Advertised ASIN", "Advertised SKU ASIN", "advertisedAsin"],
    "asin": ["(Child) ASIN", "Child ASIN", "childAsin", "ASIN"],
    "parent_asin": ["(Parent) ASIN", "Parent ASIN", "parentAsin"],
    "sku": ["SKU", "Advertised SKU", "advertisedSku", "Merchant SKU"],
    "date": ["Date", "Report Date"],
    "start_date": ["Start Date", "startDate"],
    "end_date": ["End Date", "endDate"],
    "currency": ["Currency", "Currency Code", "currencyCode"],
    "marketplace": ["Marketplace", "Country", "Marketplace Name"],
    "impressions": ["Impressions"], "clicks": ["Clicks"],
    "spend": ["Spend", "Cost", "Total Cost"],
    "sales": ["Sales", "Total Sales", "Attributed Sales", "7 Day Total Sales", "14 Day Total Sales", "30 Day Total Sales", "sales7d", "sales14d", "sales30d"],
    "orders": ["Orders", "Purchases", "Total Orders", "7 Day Total Orders (#)", "14 Day Total Orders (#)", "30 Day Total Orders (#)", "7 Day Total Purchases", "14 Day Total Purchases", "purchases7d", "purchases14d", "purchases30d"],
    "units": ["Units", "7 Day Total Units (#)", "14 Day Total Units (#)", "unitsSoldClicks7d", "unitsSoldClicks14d"],
    "sessions": ["Sessions", "Sessions - Total"],
    "ordered_sales": ["Ordered Product Sales"],
    "ordered_units": ["Units Ordered"],
    "order_items": ["Total Order Items"],
    "entity": ["Entity", "Record Type"],
    "operation": ["Operation"], "state": ["State", "Status"],
    "bid": ["Bid", "Keyword Bid"],
    "budget": ["Daily Budget", "Campaign Daily Budget"],
    "product": ["Product", "Ad Product", "adProduct"],
}
NORMAL_FIELDS = list(ALIASES) + ["resolved_asin", "asin_evidence"]
CONTEXT_FIELDS = ["account", "marketplace", "currency", "timezone", "start_date", "end_date", "aggregation", "date_basis", "attribution_days", "attribution_model", "sales_scope", "scope"]
SPEC_FIELDS = set(CONTEXT_FIELDS) | {"id", "path", "type", "ad_type", "sheet", "header_row", "column_map", "number_locale", "date_format", "encoding", "delimiter", "exclude_rows", "confirmed_data_rows"}


class InputError(ValueError):
    pass


def norm(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def text_cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def raw_cell(value):
    return "" if value is None else value.isoformat() if isinstance(value, (date, datetime)) else str(value)


def parse_date(value, fmt="iso"):
    s = text_cell(value)
    formats = {"iso": "%Y-%m-%d", "mdy": "%m/%d/%Y", "dmy": "%d/%m/%Y"}
    if fmt not in formats:
        raise InputError("date_format must be iso, mdy or dmy")
    try:
        if fmt == "iso" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
            raise ValueError()
        return datetime.strptime(s, formats[fmt]).date().isoformat()
    except ValueError as e:
        raise InputError(f"Invalid {fmt} calendar date: {s}") from e


def number(value, locale="dot", count=False, optional=False):
    s = text_cell(value)
    if not s:
        if optional:
            return None
        raise InputError("Missing numeric value; blank is not zero")
    if len(s) > 60:
        raise InputError("Numeric cell is too long")
    if isinstance(value, bool) or "%" in s or re.search(r"[eE]", s):
        raise InputError("Expected a number, not a percentage, boolean or exponent")
    # Strip only known monetary markers. Their presence does not establish currency.
    s = re.sub(r"^(?:USD|CAD|GBP|US\$|CA\$|C\$|\$|£)\s*", "", s, flags=re.I)
    s = s.replace("\u00a0", " ").strip()
    if locale == "dot":
        pattern = r"(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?"
        cleaned = s.replace(",", "")
    elif locale == "comma":
        pattern = r"(?:\d+|\d{1,3}(?:[. ]\d{3})+)(?:,\d+)?"
        cleaned = s.replace(".", "").replace(" ", "").replace(",", ".")
    else:
        raise InputError("number_locale must be dot or comma")
    if not re.fullmatch(pattern, s):
        raise InputError(f"Invalid non-negative {locale}-locale number: {s}")
    try:
        n = Decimal(cleaned)
    except InvalidOperation as e:
        raise InputError("Invalid numeric value") from e
    if not n.is_finite() or n > Decimal("1000000000000"):
        raise InputError("Number exceeds the supported 1 trillion limit")
    if count and n != n.to_integral_value():
        raise InputError("Count must be a whole number")
    if -n.as_tuple().exponent > 12:
        raise InputError("At most 12 decimal places are supported")
    return str(int(n)) if count else format(n, "f")


def market(value):
    s = text_cell(value).upper()
    return {"GB": "UK", "UNITED KINGDOM": "UK", "AMAZON.CO.UK": "UK", "UNITED STATES": "US", "AMAZON.COM": "US", "CANADA": "CA", "AMAZON.CA": "CA"}.get(s, s)


def ad_product(value):
    s = norm(value)
    return {"sp": "SP", "sponsoredproducts": "SP", "sb": "SB", "sponsoredbrands": "SB", "sd": "SD", "sponsoreddisplay": "SD"}.get(s)


def window(header):
    s = norm(header)
    m = re.search(r"(\d+)day", s) or re.search(r"(?:sales|purchases|unitsSoldClicks)(\d+)d$", str(header), re.I)
    return int(m.group(1)) if m else None


def metric_scope(header):
    h = norm(header)
    if "advertisedsku" in h or "advertisedasin" in h or "samesku" in h:
        return "advertised_sku"
    if "othersku" in h or "otherasin" in h:
        return "other_sku"
    if "total" in h:
        return "total"
    return None


def issue(report, severity, code, message, row=None, field=None):
    finding = {"severity": severity, "code": code, "message": message}
    if row is not None:
        finding["source_row"] = row
    if field:
        finding["field"] = field
    report["findings"].append(finding)


def read_tables(path, spec):
    if not path.is_file():
        raise InputError("Source file does not exist")
    if path.stat().st_size > MAX_BYTES:
        raise InputError("Source exceeds the 100 MiB input limit")
    ext = path.suffix.lower()
    if ext == ".xlsx":
        try:
            import openpyxl
        except ImportError as e:
            raise InputError("XLSX needs openpyxl: python -m pip install openpyxl. CSV/TSV need no package.") from e
        with zipfile.ZipFile(path) as z:
            if sum(x.file_size for x in z.infolist()) > 512 * 1024 * 1024:
                raise InputError("XLSX expanded content exceeds 512 MiB")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if spec.get("sheet") and spec["sheet"] not in wb.sheetnames:
                raise InputError("Requested XLSX sheet is missing")
            sheets = [wb[spec["sheet"]]] if spec.get("sheet") else wb.worksheets
            for ws in sheets:
                if ws.max_column and ws.max_column > 500:
                    raise InputError("Sheet exceeds the supported 500 column limit")
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                    if i > MAX_ROWS + 100:
                        raise InputError("Sheet exceeds the 250,000 row limit")
                    rows.append((i, list(row)))
                if any(any(text_cell(c) for c in row) for _, row in rows):
                    yield ws.title, rows
        finally:
            wb.close()
    elif ext in {".csv", ".tsv", ".txt"}:
        encoding = spec.get("encoding", "utf-8-sig")
        if encoding not in {"utf-8-sig", "utf-8", "utf-16", "cp1252"}:
            raise InputError("Unsupported encoding")
        content = path.read_text(encoding=encoding)
        delimiter = spec.get("delimiter") or ("\t" if ext == ".tsv" else None)
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(content[:65536], delimiters=",\t;").delimiter
            except csv.Error:
                raise InputError("Cannot establish delimiter; declare delimiter in the manifest")
        if delimiter not in {",", "\t", ";"}:
            raise InputError("delimiter must be comma, tab or semicolon")
        csv.field_size_limit(1024 * 1024)
        reader = csv.reader(io.StringIO(content, newline=""), delimiter=delimiter, strict=True)
        rows = []
        for i, row in enumerate(reader, 1):
            if i > MAX_ROWS + 100:
                raise InputError("File exceeds the 250,000 row limit")
            rows.append((i, row))
        yield "", rows
    else:
        raise InputError("Supported formats are CSV, TSV, TXT and XLSX; legacy XLS is unsupported")


def detect(headers):
    h = {norm(x) for x in headers}
    has = lambda field: bool(h & {norm(x) for x in ALIASES[field]})
    if has("entity"):
        return "bulk"
    if has("ordered_sales"):
        return "business"
    if has("search_term"):
        return "search_term"
    if has("advertised_asin") or (has("sku") and has("sales")):
        return "advertised_product"
    if has("target"):
        return "targeting"
    if has("campaign_id") or has("campaign_name"):
        return "campaign"
    return None


def header_index(rows, spec):
    explicit = spec.get("header_row")
    if explicit is not None:
        if not isinstance(explicit, int) or isinstance(explicit, bool) or explicit < 1:
            raise InputError("header_row must be a positive physical row number")
        for i, (line, _) in enumerate(rows):
            if line == explicit:
                return i
        raise InputError("Requested header row is missing")
    for i, (_, row) in enumerate(rows[:50]):
        if detect(row):
            return i
    raise InputError("No recognized English report header; supply header_row, type and column_map")


def mapping(headers, spec, report):
    result = {}
    supplied = spec.get("column_map", {})
    if not isinstance(supplied, dict) or set(supplied) - set(ALIASES):
        raise InputError("column_map must map supported canonical fields to exact header text")
    for field, aliases in ALIASES.items():
        if field in supplied:
            found = [i for i, h in enumerate(headers) if h == supplied[field]]
        else:
            found = [i for i, h in enumerate(headers) if norm(h) in {norm(a) for a in aliases}]
        if len(found) > 1:
            issue(report, "block", "AMBIGUOUS_COLUMN", f"Multiple columns map to {field}; select one with column_map", field=field)
        elif found:
            result[field] = found[0]
        elif field in supplied:
            issue(report, "block", "MAPPED_COLUMN_MISSING", f"Declared header for {field} is absent", field=field)
    if len(set(result.values())) != len(result):
        issue(report, "block", "COLUMN_REUSED", "One source column maps to several canonical fields")
    return result


def require_columns(kind, mapping_, report):
    required = {"spend", "clicks", "sales", "orders"} if kind in {"campaign", "search_term", "targeting", "advertised_product"} else {"ordered_sales"} if kind == "business" else {"entity"}
    alternatives = {"campaign": [("campaign_id", "campaign_name")], "search_term": [("campaign_id", "campaign_name"), ("search_term",)], "targeting": [("campaign_id", "campaign_name"), ("target", "keyword_id", "target_id")], "advertised_product": [("campaign_id", "campaign_name"), ("advertised_asin", "sku")], "business": [("asin", "sku")], "bulk": [("campaign_id", "campaign_name")]}
    for field in sorted(required - set(mapping_)):
        issue(report, "block", "REQUIRED_COLUMN_MISSING", f"{kind} requires {field}; use a matching report or explicit mapping", field=field)
    for group in alternatives.get(kind, []):
        if not set(group) & set(mapping_):
            issue(report, "block", "DIMENSION_MISSING", "Missing identifying column: " + " or ".join(group))


def validate_context(spec, report, as_of):
    c = report["context"]
    if c.get("account") is not None and (not isinstance(c["account"], str) or not c["account"].strip()):
        issue(report, "block", "INVALID_ACCOUNT", "account must be a nonempty identifier string")
    ads = report["type"] not in {"business", "bulk"}
    required = ["account", "marketplace", "currency", "timezone", "start_date", "end_date", "aggregation", "scope"]
    if ads:
        required += ["ad_type", "date_basis", "attribution_days", "attribution_model", "sales_scope"]
    elif report["type"] == "business":
        required += ["date_basis"]
    # Bulk exports contain different entity grains and may have no performance window.
    if report["type"] == "bulk":
        required = ["account", "marketplace", "currency"]
    for field in required:
        if c.get(field) is None or c.get(field) in ("", "unknown"):
            issue(report, "context", "CONTEXT_MISSING", f"Confirm {field}; do not infer it from a filename", field=field)
    if c.get("marketplace"):
        c["marketplace"] = market(c["marketplace"])
        if c["marketplace"] not in MARKETS:
            issue(report, "block", "UNSUPPORTED_MARKETPLACE", "First release supports US, UK and CA")
    if c.get("currency"):
        c["currency"] = str(c["currency"]).upper()
        if c.get("marketplace") in MARKETS and c["currency"] != MARKETS[c["marketplace"]]:
            issue(report, "block", "MARKET_CURRENCY_CONFLICT", "Currency does not match the declared marketplace; no FX conversion is applied")
    if c.get("timezone"):
        try:
            ZoneInfo(c["timezone"])
        except (ZoneInfoNotFoundError, ValueError, TypeError):
            issue(report, "block", "INVALID_TIMEZONE", "Use the actual report timezone as an IANA name")
    for key in ["start_date", "end_date"]:
        if c.get(key):
            try:
                c[key] = parse_date(c[key])
            except InputError as e:
                issue(report, "block", "INVALID_PERIOD", str(e), field=key)
    start, end = c.get("start_date"), c.get("end_date")
    if start and end and re.fullmatch(r"\d{4}-\d{2}-\d{2}", start) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", end):
        try:
            if date.fromisoformat(start) > date.fromisoformat(end):
                issue(report, "block", "REVERSED_PERIOD", "start_date is after end_date")
            if date.fromisoformat(end) > as_of:
                issue(report, "block", "FUTURE_PERIOD", "end_date is later than the declared as-of date")
            days = c.get("attribution_days")
            if ads and isinstance(days, int) and not isinstance(days, bool) and days > 0 and as_of <= date.fromisoformat(end) + timedelta(days=days):
                issue(report, "context", "ATTRIBUTION_UNSETTLED", "Full attribution window has not elapsed; exported metrics may still change")
        except ValueError:
            pass
    for key, allowed in {"aggregation": {"daily", "summary"}, "ad_type": {"SP", "SB", "SD"}, "date_basis": {"ad_interaction", "order", "conversion"}, "attribution_model": {"click", "click_and_view"}, "sales_scope": {"total", "advertised_sku", "other_sku"}}.items():
        if c.get(key) not in (None, "", "unknown") and c[key] not in allowed:
            issue(report, "block", "INVALID_CONTEXT", f"Unsupported {key}", field=key)
    if c.get("attribution_days") is not None and (not isinstance(c["attribution_days"], int) or isinstance(c["attribution_days"], bool) or not 1 <= c["attribution_days"] <= 90):
        issue(report, "block", "INVALID_ATTRIBUTION_WINDOW", "attribution_days must be an integer from 1 to 90")
    if c.get("scope") is not None and (not isinstance(c["scope"], str) or not c["scope"].strip()):
        issue(report, "block", "INVALID_SCOPE", "scope must describe the account/product/filter coverage")


def inspect_table(path, sheet, raw_rows, spec, as_of, dedupe):
    report = {"id": spec["id"] + (":" + sheet if sheet else ""), "source_file": str(path), "sheet": sheet, "type": spec.get("type", "auto"), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "findings": [], "decisions": [], "rows": [], "context": {k: spec.get(k) for k in CONTEXT_FIELDS + ["ad_type"]}, "totals": None}
    hi = header_index(raw_rows, spec)
    headers = [text_cell(x) for x in raw_rows[hi][1]]
    if len(headers) > 500:
        raise InputError("Report exceeds the supported 500 column limit")
    # Excel exports may have styled trailing columns; trim only fully empty tails.
    while headers and headers[-1] == "" and all(len(row) <= len(headers)-1 or not text_cell(row[len(headers)-1]) for _, row in raw_rows[hi+1:]):
        headers.pop()
    report["headers"] = headers
    report["header_row"] = raw_rows[hi][0]
    if hi:
        report["decisions"].append({"action": "skip_preamble", "rows": [line for line, _ in raw_rows[:hi]], "reason": "Rows precede the detected or declared header; raw source remains unchanged"})
    if any(not h for h in headers) or len(set(map(norm, headers))) != len(headers):
        issue(report, "block", "DUPLICATE_OR_BLANK_HEADER", "Headers must be nonempty and unique after punctuation/case normalization")
    detected = detect(headers)
    if report["type"] == "auto":
        report["type"] = detected
    elif detected and detected != report["type"]:
        issue(report, "block", "REPORT_TYPE_CONFLICT", f"Declared {report['type']} but headers resemble {detected}; select the right sheet/report")
    if report["type"] not in KINDS - {"auto"}:
        raise InputError("Cannot determine report type; declare a supported type and column_map")
    mp = mapping(headers, spec, report)
    report["column_map"] = {k: headers[i] for k, i in mp.items()}
    require_columns(report["type"], mp, report)
    # A named conversion window in the file is evidence, not a marketplace default.
    windows = {field: window(headers[mp[field]]) for field in ["sales", "orders", "units"] if field in mp}
    known = {x for x in windows.values() if x}
    report["metric_windows"] = windows
    if len(known) > 1:
        issue(report, "block", "METRIC_WINDOW_CONFLICT", "Purchases, units and sales use different attribution windows")
    elif known:
        detected_window = next(iter(known))
        if report["context"].get("attribution_days") is not None and report["context"]["attribution_days"] != detected_window:
            issue(report, "block", "DECLARED_WINDOW_CONFLICT", "Header attribution window differs from the manifest")
        else:
            report["context"]["attribution_days"] = detected_window
    scopes = {field: metric_scope(headers[mp[field]]) for field in ["sales", "orders", "units"] if field in mp}
    report["metric_scopes"] = scopes
    known_scopes = {s for s in scopes.values() if s}
    if len(known_scopes) > 1:
        issue(report, "block", "METRIC_SCOPE_CONFLICT", "Purchases, units and sales mix total and advertised/other SKU definitions")
    if scopes.get("sales") and report["context"].get("sales_scope") and scopes["sales"] != report["context"]["sales_scope"]:
        issue(report, "block", "SALES_SCOPE_CONFLICT", "Sales column definition conflicts with the declared sales_scope")
    observed_currencies, observed_markets, observed_dates = set(), set(), set()
    exact = defaultdict(list)
    excludes = spec.get("exclude_rows", [])
    confirmed = spec.get("confirmed_data_rows", [])
    if not isinstance(excludes, list) or any(not isinstance(x, int) or isinstance(x, bool) for x in excludes):
        raise InputError("exclude_rows must contain physical row numbers")
    if not isinstance(confirmed, list) or any(not isinstance(x, int) or isinstance(x, bool) for x in confirmed) or set(confirmed) & set(excludes):
        raise InputError("confirmed_data_rows must contain physical row numbers distinct from exclude_rows")
    available = {x for x, _ in raw_rows[hi+1:]}
    if set(excludes) - available:
        issue(report, "block", "EXCLUDED_ROW_MISSING", "An excluded physical row does not exist in the selected table")
    if set(confirmed) - available:
        issue(report, "block", "CONFIRMED_ROW_MISSING", "A confirmed physical data row does not exist")
    for line, values in raw_rows[hi+1:]:
        if not any(text_cell(x) for x in values):
            report["decisions"].append({"action": "skip_blank_row", "source_row": line})
            continue
        if line in excludes:
            report["decisions"].append({"action": "exclude_declared_row", "source_row": line, "raw_values": list(map(raw_cell, values))})
            continue
        if line in confirmed:
            report["decisions"].append({"action": "confirm_data_row", "source_row": line, "reason": "Caller explicitly confirmed this is data, not a summary row"})
        if len(values) > len(headers) and all(not text_cell(x) for x in values[len(headers):]):
            values = values[:len(headers)]
        elif path.suffix.lower() == ".xlsx" and len(values) < len(headers):
            values += [None] * (len(headers) - len(values))
        row = {"source_row": line, "values": {}, "raw_values": list(map(raw_cell, values)), "issue_codes": [], "excluded_duplicate": False}
        before = len(report["findings"])
        if len(values) != len(headers):
            issue(report, "block", "ROW_WIDTH", "Column count differs from the header; no fields were shifted", line)
        else:
            exact[tuple(row["raw_values"])].append(row)
            for field, i in mp.items():
                value, s = values[i], text_cell(values[i])
                try:
                    if s.startswith("=") and path.suffix.lower() == ".xlsx":
                        raise InputError("Formula cells are not evaluated; provide a values-only export")
                    if field in METRICS:
                        optional = field in {"bid", "budget", "units", "impressions", "sessions", "ordered_units", "order_items"} or report["type"] == "bulk"
                        row["values"][field] = number(value, spec.get("number_locale", "dot"), field in COUNTS, optional)
                    elif field in IDS:
                        if isinstance(value, (float, int)) and (len(str(int(value))) > 15 or isinstance(value, float) and not value.is_integer()):
                            raise InputError("Numeric identifier may have lost Excel precision; re-export IDs as text")
                        if isinstance(value, (float, int)) and not isinstance(value, bool):
                            s = str(int(value))
                        if s and not re.fullmatch(r"\d+", s):
                            raise InputError("Identifier must be digit text, without decimals or scientific notation")
                        row["values"][field] = s
                    elif field in {"date", "start_date", "end_date"}:
                        if s:
                            if isinstance(value, (datetime, date)):
                                s = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                                row["values"][field] = parse_date(s)
                            else:
                                row["values"][field] = parse_date(s, spec.get("date_format", "iso"))
                        else:
                            raise InputError("Missing date in a mapped date column")
                    elif field in {"asin", "parent_asin", "advertised_asin"}:
                        if s and not re.fullmatch(r"[A-Za-z0-9]{10}", s):
                            raise InputError("ASIN must contain exactly 10 letters/digits")
                        row["values"][field] = s.upper()
                    else:
                        row["values"][field] = s
                except (InputError, ValueError, OverflowError) as e:
                    issue(report, "block", "INVALID_CELL", str(e), line, field)
            v = row["values"]
            alternatives = {"campaign": [("campaign_id", "campaign_name")], "search_term": [("campaign_id", "campaign_name"), ("search_term",)], "targeting": [("campaign_id", "campaign_name"), ("target", "keyword_id", "target_id")], "advertised_product": [("campaign_id", "campaign_name"), ("advertised_asin", "sku")], "business": [("asin", "sku")], "bulk": [("entity",)]}
            for group in alternatives[report["type"]]:
                if not any(v.get(k) for k in group):
                    issue(report, "block", "ROW_DIMENSION_MISSING", "Row lacks identifying value: " + " or ".join(group), line)
            for field in ["currency", "marketplace"]:
                if field in mp:
                    if not v.get(field):
                        issue(report, "block", "EMPTY_SCOPE_CELL", f"Mapped {field} is blank", line, field)
                    elif field == "currency":
                        observed_currencies.add(v[field].upper())
                    else:
                        observed_markets.add(market(v[field]))
            # Explicit currency codes in amount cells/headers are also evidence.
            for field in MONEY & set(mp):
                for value in [text_cell(values[mp[field]]), headers[mp[field]]]:
                    match = re.search(r"\b(USD|CAD|GBP)\b|(?:US\$|CA\$|C\$)", value, re.I)
                    if match:
                        observed_currencies.add({"US$": "USD", "CA$": "CAD", "C$": "CAD"}.get(match.group(0).upper(), match.group(0).upper()))
                    if "£" in value:
                        observed_currencies.add("GBP")
            if v.get("product") and report["context"].get("ad_type") and ad_product(v["product"]) != report["context"]["ad_type"]:
                issue(report, "block", "AD_PRODUCT_CONFLICT", "Row ad product differs from declared ad_type", line)
            if v.get("date"):
                observed_dates.add(v["date"])
                for key, op in [("start_date", lambda a, b: a < b), ("end_date", lambda a, b: a > b)]:
                    if report["context"].get(key) and op(v["date"], report["context"][key]):
                        issue(report, "block", "DATE_OUTSIDE_PERIOD", "Daily date is outside the declared report period", line)
            for field in ["start_date", "end_date"]:
                if v.get(field) and report["context"].get(field) and v[field] != report["context"][field]:
                    issue(report, "block", "ROW_PERIOD_CONFLICT", f"Row {field} differs from the declared period", line)
            for field in ["campaign_name", "search_term", "target", "asin", "sku"]:
                if line not in confirmed and norm(v.get(field, "")) in {"total", "totals", "grandtotal"}:
                    issue(report, "context", "POSSIBLE_TOTAL_ROW", "Confirm whether this is a summary row or a real entity; use exclude_rows only after inspection", line, field)
            row["issue_codes"] = sorted({x["code"] for x in report["findings"][before:]})
        report["rows"].append(row)
    if not report["rows"]:
        issue(report, "block", "NO_DATA", "No data rows remain")
    for field, observed in [("currency", observed_currencies), ("marketplace", observed_markets)]:
        if len(observed) > 1:
            issue(report, "block", "MIXED_" + field.upper(), f"Mixed {field} values; no amounts were pooled")
        elif observed:
            value = next(iter(observed))
            if report["context"].get(field) and (market(report["context"][field]) if field == "marketplace" else str(report["context"][field]).upper()) != value:
                issue(report, "block", "DECLARED_" + field.upper() + "_CONFLICT", f"Source {field} differs from the manifest")
            elif not report["context"].get(field):
                report["context"][field] = value
    report["observed_daily_dates"] = sorted(observed_dates)
    for group in exact.values():
        if len(group) > 1:
            lines = [r["source_row"] for r in group]
            if dedupe:
                for row in group[1:]:
                    row["excluded_duplicate"] = True
                report["decisions"].append({"action": "dedupe_exact", "kept_row": lines[0], "excluded_rows": lines[1:]})
                issue(report, "warning", "EXACT_DUPLICATES_REMOVED", f"Explicit dedupe retained the first exact row from {lines}")
            else:
                issue(report, "block", "EXACT_DUPLICATES", f"Exact repeated rows at {lines}; inspect before explicitly deduplicating")
                for row in group:
                    row["issue_codes"].append("EXACT_DUPLICATES")
    validate_context(spec, report, as_of)
    if report["context"].get("aggregation") == "summary" and "date" in mp:
        issue(report, "context", "AGGREGATION_REVIEW", "Summary declaration includes a daily Date column; confirm aggregation rather than pooling daily and summary rows")
    if report["type"] != "bulk" and report["context"].get("aggregation") == "daily" and "date" not in mp:
        issue(report, "block", "DAILY_DATE_MISSING", "Declared daily report has no mapped daily Date column")
    if report["type"] in {"campaign", "search_term", "targeting", "advertised_product"} and "campaign_id" not in mp:
        issue(report, "warning", "CAMPAIGN_ID_UNAVAILABLE", "Name-based analysis remains possible; ID-based joins and account-change preparation are unavailable")
    report["decisions"].append({"action": "preserve_unmapped_columns", "headers": [h for i, h in enumerate(headers) if i not in mp.values()]})
    return report


def finalize(report):
    levels = {x["severity"] for x in report["findings"]}
    report["status"] = "BLOCKED" if "block" in levels else "NEEDS_CONTEXT" if "context" in levels else "READY"
    for row in report["rows"]:
        row["status"] = "EXCLUDED_DUPLICATE" if row["excluded_duplicate"] else report["status"]
    if report["status"] == "READY" and report["type"] != "bulk":
        totals = {}
        for field in sorted(METRICS & set(report.get("column_map", {}))):
            values = [r["values"].get(field) for r in report["rows"] if not r["excluded_duplicate"]]
            totals[field] = None if any(v is None for v in values) else format(sum((Decimal(v) for v in values), Decimal(0)), "f")
        report["totals"] = totals
    else:
        report["totals"] = None


def join_checks(reports):
    catalogs, bulk_map = defaultdict(set), defaultdict(set)
    for report in reports:
        if report["status"] != "READY":
            continue
        c = report["context"]
        scope = (c.get("account"), c.get("marketplace"))
        for row in report["rows"]:
            if row["excluded_duplicate"]:
                continue
            v = row["values"]
            if report["type"] == "business" and v.get("asin"):
                catalogs[scope].add(v["asin"])
            if report["type"] == "bulk" and norm(v.get("entity", "")) in {"productad", "ad"} and v.get("campaign_id") and v.get("ad_group_id"):
                asin = v.get("advertised_asin") or v.get("asin")
                if asin:
                    bulk_map[scope + (v["campaign_id"], v["ad_group_id"])].add(asin)
    results = []
    for report in reports:
        if report["status"] != "READY" or report["type"] in {"bulk", "business"}:
            continue
        c = report["context"]
        scope = (c.get("account"), c.get("marketplace"))
        counts = defaultdict(int)
        for row in report["rows"]:
            if row["excluded_duplicate"]:
                continue
            v = row["values"]
            asin = v.get("advertised_asin")
            if asin:
                v["asin_evidence"] = "source advertised_asin"
            elif report["type"] in {"search_term", "targeting"} and v.get("campaign_id") and v.get("ad_group_id"):
                candidates = bulk_map.get(scope + (v["campaign_id"], v["ad_group_id"]), set())
                if len(candidates) == 1:
                    asin = next(iter(candidates))
                    v["asin_evidence"] = "unique bulk ad-group product mapping; current configuration, not historical attribution"
                elif len(candidates) > 1:
                    counts["ambiguous_ad_group"] += 1
            if asin:
                v["resolved_asin"] = asin
                counts["resolved"] += 1
                if catalogs.get(scope) and asin not in catalogs[scope]:
                    counts["not_in_supplied_business_report"] += 1
            else:
                counts["unresolved"] += 1
        results.append({"report_id": report["id"], "counts": dict(counts), "note": "Configuration joins do not allocate ASIN sales, prove historical product coverage or identify purchased-product halo. Missing Business Report ASINs can reflect scope or inactivity."})
    return results


def comparisons(reports):
    results = []
    for i, left in enumerate(reports):
        for right in reports[i+1:]:
            if {left["type"], right["type"]} == {"bulk"} or "bulk" in {left["type"], right["type"]}:
                continue
            fields = ["account", "marketplace", "currency", "timezone", "start_date", "end_date", "scope", "date_basis"]
            if left["type"] != "business" and right["type"] != "business":
                fields += ["ad_type", "attribution_days", "attribution_model", "sales_scope"]
            mismatches = [k for k in fields if left["context"].get(k) != right["context"].get(k)]
            missing = [k for k in fields if left["context"].get(k) in (None, "", "unknown") or right["context"].get(k) in (None, "", "unknown")]
            compatible = left["status"] == right["status"] == "READY" and not mismatches and not missing
            result = {"left": left["id"], "right": right["id"], "status": "BASIS_ALIGNED" if compatible else "NOT_COMPARABLE", "mismatched_fields": mismatches, "unverified_fields": missing, "metric_differences": None}
            if left["type"] == "business" or right["type"] == "business":
                result["status"] = "DEFINITION_REVIEW"
                result["note"] = "Ad-attributed sales and ordered product sales are different measures. No exact organic-sales subtraction or order-level reconciliation is produced."
            elif compatible:
                shared = set(left["totals"] or {}) & set(right["totals"] or {})
                result["metric_differences"] = {k: format(Decimal(left["totals"][k]) - Decimal(right["totals"][k]), "f") for k in sorted(shared) if left["totals"][k] is not None and right["totals"][k] is not None}
                result["note"] = "Left minus right diagnostic totals, not an expected zero reconciliation. Report coverage and grain may differ; search-term impressions cover clicked queries only."
            results.append(result)
    return results


def safe_csv(value):
    s = "" if value is None else str(value)
    return "'" + s if s.lstrip().startswith(("=", "+", "-", "@")) else s


def safe_md(value):
    return str(value).replace("\\", "\\\\").replace("|", "\\|").replace("\n", " ").replace("\r", " ")[:400]


def write_outputs(out, result):
    (out / "report.json").write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (out / "comparisons.json").write_text(json.dumps(result["comparisons"], indent=2) + "\n", encoding="utf-8")
    (out / "joins.json").write_text(json.dumps(result["joins"], indent=2) + "\n", encoding="utf-8")
    meta = ["report_id", "report_type", "source_file", "sheet", "source_row", "row_status", "account", "marketplace", "currency", "report_start_date", "report_end_date", "issue_codes"]
    fields = meta + [k for k in NORMAL_FIELDS if k not in meta] + ["raw_json"]
    for filename, ready_only in [("normalized-audit.csv", False), ("analysis-ready.csv", True)]:
        with (out / filename).open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fields)
            writer.writeheader()
            for report in result["reports"]:
                for row in report["rows"]:
                    if ready_only and row["status"] != "READY":
                        continue
                    c = report["context"]
                    r = {"report_id": report["id"], "report_type": report["type"], "source_file": report["source_file"], "sheet": report["sheet"], "source_row": row["source_row"], "row_status": row["status"], "account": c.get("account"), "marketplace": c.get("marketplace"), "currency": c.get("currency"), "report_start_date": c.get("start_date"), "report_end_date": c.get("end_date"), "issue_codes": ";".join(row["issue_codes"]), "raw_json": json.dumps({"headers": report.get("headers", []), "values": row["raw_values"]}, ensure_ascii=False), **row["values"]}
                    # Row currency/marketplace cannot overwrite authoritative normalized context.
                    r["currency"], r["marketplace"] = c.get("currency"), c.get("marketplace")
                    writer.writerow({k: safe_csv(r.get(k)) for k in fields})
    with (out / "exceptions.csv").open("w", newline="", encoding="utf-8") as f:
        fields_ = ["report_id", "severity", "code", "source_row", "field", "message"]
        writer = csv.DictWriter(f, fields_)
        writer.writeheader()
        for report in result["reports"]:
            for finding in report["findings"]:
                writer.writerow({k: safe_csv((finding | {"report_id": report["id"]}).get(k)) for k in fields_})
        for finding in result["batch_findings"]:
            writer.writerow({k: safe_csv((finding | {"report_id": "BATCH"}).get(k)) for k in fields_})
    md = ["# Amazon Report Preflight", "", f"Status: **{result['status']}** · Method {VERSION} · As of {result['as_of']}", "", "Sources were read only. Analysis-ready rows remain separated by report; do not pool reports of different types or overlapping scopes.", "", "| Report | Type | Status | Rows retained |", "| --- | --- | --- | --- |"]
    for report in result["reports"]:
        md.append(f"| {safe_md(report['id'])} | {report['type']} | {report['status']} | {len(report['rows'])} |")
    md += ["", "## Findings", ""]
    for report in result["reports"]:
        md.append("### " + safe_md(report["id"]))
        for finding in report["findings"]:
            location = f" (source row {finding['source_row']})" if finding.get("source_row") else ""
            md.append(f"- {finding['severity']}: {finding['code']}{location} — {safe_md(finding['message'])}")
        if not report["findings"]:
            md.append("- No issues in the checked rules.")
        md.append("")
    for f in result["batch_findings"]:
        md.append(f"- Batch {f['severity']}: {f['code']} — {safe_md(f['message'])}")
    md += ["", "## Limits", "", "- Unknown context is not a pass. A blocked or context-incomplete report contributes no analysis-ready rows or totals.", "- Exact duplicates are retained and block analysis unless --dedupe-exact was explicitly chosen. All exclusions are recorded.", "- CSV text that could execute as a spreadsheet formula is prefixed with an apostrophe. Original text remains in raw_json and report.json.", "- Current bulk ASIN configuration is not historical attribution. Mixed-ASIN ad groups remain unresolved.", "- Ad-attributed and ordered product sales are not equivalent. No exact organic revenue, profit, causal effect or account-change recommendation is inferred.", "- READY means the checked export and declared context are usable, not that all account data, definitions or attribution adjustments were independently verified."]
    (out / "report.md").write_text("\n".join(md) + "\n", encoding="utf-8")


def validate_manifest(manifest):
    if not isinstance(manifest, dict) or set(manifest) - {"version", "as_of", "reports", "required_reports"}:
        raise InputError("Manifest must contain only version, as_of, reports and optional required_reports")
    if manifest.get("version") != 1 or isinstance(manifest.get("version"), bool):
        raise InputError("Manifest version must be 1")
    specs = manifest.get("reports")
    if not isinstance(specs, list) or not specs:
        raise InputError("Manifest needs at least one report")
    seen = set()
    for spec in specs:
        if not isinstance(spec, dict) or set(spec) - SPEC_FIELDS:
            raise InputError("Report contains unsupported keys; check the manifest reference")
        if not isinstance(spec.get("id"), str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,64}", spec["id"]):
            raise InputError("Each report needs a unique simple id")
        if spec["id"] in seen:
            raise InputError("Duplicate report id")
        seen.add(spec["id"])
        if not isinstance(spec.get("path"), str) or not spec["path"]:
            raise InputError("Each report needs a path")
        if spec.get("type", "auto") not in KINDS:
            raise InputError("Unsupported report type")
        if spec.get("number_locale", "dot") not in {"dot", "comma"} or spec.get("date_format", "iso") not in {"iso", "mdy", "dmy"}:
            raise InputError("Unsupported number_locale or date_format")
    required = manifest.get("required_reports", [])
    if not isinstance(required, list) or any(x not in KINDS - {"auto"} for x in required):
        raise InputError("required_reports must list supported report types")
    return specs


def run(manifest, base_dir, out, as_of=None, dedupe=False):
    specs = validate_manifest(manifest)
    as_of = date.fromisoformat(parse_date(as_of or manifest.get("as_of") or datetime.now().date().isoformat()))
    out = Path(out).resolve()
    if out.exists():
        raise InputError("Output directory already exists; choose a new directory")
    source_paths = [(Path(spec["path"]) if Path(spec["path"]).is_absolute() else Path(base_dir) / spec["path"]).resolve() for spec in specs]
    if any(out == path or out in path.parents for path in source_paths):
        raise InputError("Output directory cannot contain a source file")
    reports, batch = [], []
    seen_sources = set()
    for spec, path in zip(specs, source_paths):
        added = 0
        try:
            for sheet, rows in read_tables(path, spec):
                try:
                    report = inspect_table(path, sheet, rows, spec, as_of, dedupe)
                except (InputError, ValueError, csv.Error, TypeError) as e:
                    report = {"id": spec["id"] + (":" + sheet if sheet else ""), "source_file": str(path), "sheet": sheet, "type": spec.get("type", "auto"), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "findings": [{"severity": "block", "code": "TABLE_INPUT_ERROR", "message": str(e)}], "decisions": [], "rows": [], "context": {k: spec.get(k) for k in CONTEXT_FIELDS + ["ad_type"]}}
                source_key = (report["sha256"], sheet)
                if source_key in seen_sources:
                    issue(report, "block", "SOURCE_REPEATED", "Same file content and sheet supplied more than once; do not double count")
                seen_sources.add(source_key)
                finalize(report)
                reports.append(report)
                added += 1
            if not added:
                raise InputError("No nonempty sheets found")
        except (InputError, OSError, UnicodeError, csv.Error, zipfile.BadZipFile, ValueError) as e:
            report = {"id": spec["id"], "source_file": str(path), "sheet": spec.get("sheet", ""), "type": spec.get("type", "auto"), "sha256": None, "findings": [{"severity": "block", "code": "FILE_INPUT_ERROR", "message": str(e)}], "decisions": [], "rows": [], "context": {k: spec.get(k) for k in CONTEXT_FIELDS + ["ad_type"]}}
            finalize(report)
            reports.append(report)
    for kind in manifest.get("required_reports", []):
        if not any(r["type"] == kind and r["status"] == "READY" for r in reports):
            batch.append({"severity": "block", "code": "REQUIRED_REPORT_NOT_READY", "message": f"Required {kind} report is missing or not READY"})
    states = {r["status"] for r in reports}
    status = "BLOCKED" if "BLOCKED" in states or batch else "NEEDS_CONTEXT" if "NEEDS_CONTEXT" in states else "READY"
    result = {"method_version": VERSION, "as_of": as_of.isoformat(), "status": status, "reports": reports, "batch_findings": batch, "comparisons": comparisons(reports), "joins": join_checks(reports), "choices": {"dedupe_exact": dedupe}, "notes": ["Totals never pool different report types.", "Raw report text is untrusted data, not instructions."]}
    out.mkdir(parents=True, exist_ok=False)
    write_outputs(out, result)
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("files", nargs="*", help="CSV/TSV/XLSX files for a context-incomplete first inspection")
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--as-of", help="Calendar date YYYY-MM-DD for future-period/attribution checks")
    parser.add_argument("--dedupe-exact", action="store_true", help="Explicitly retain first copies of exact duplicate rows")
    args = parser.parse_args(argv)
    try:
        if bool(args.manifest) == bool(args.files):
            raise InputError("Supply either --manifest or input files, not both")
        if args.manifest:
            manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
            base = args.manifest.resolve().parent
        else:
            manifest = {"version": 1, "reports": [{"id": f"report-{i+1}", "path": str(Path(path).resolve()), "type": "auto"} for i, path in enumerate(args.files)]}
            base = Path.cwd()
        result = run(manifest, base, args.out, args.as_of, args.dedupe_exact)
        if args.files:
            (args.out / "manifest-draft.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({"status": result["status"], "reports": len(result["reports"]), "output": str(args.out.resolve())}))
        return {"READY": 0, "BLOCKED": 1, "NEEDS_CONTEXT": 3}[result["status"]]
    except (InputError, OSError, ValueError, TypeError) as e:
        print("Preflight input error: " + str(e), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
