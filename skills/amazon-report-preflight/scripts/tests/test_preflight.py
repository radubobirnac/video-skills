"""Synthetic regressions for normalization, scope, lineage and blocked-data behavior."""
import csv
import hashlib
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

SCRIPT = Path(__file__).resolve().parents[1] / "report_preflight.py"
spec = importlib.util.spec_from_file_location("preflight", SCRIPT)
p = importlib.util.module_from_spec(spec)
spec.loader.exec_module(p)
try:
    import openpyxl
except ImportError:
    openpyxl = None

HEAD = ["Campaign ID", "Campaign Name", "Clicks", "Spend", "7 Day Total Orders (#)", "7 Day Total Sales", "Currency"]
ROW = ["12345678901234567890", "Sample | SP - AUTO", "10", "1.10", "2", "20.20", "USD"]


class PreflightTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.n = 0

    def tearDown(self):
        self.temp.cleanup()

    def file(self, name="campaign.csv", headers=None, rows=None, delimiter=",", encoding="utf-8", preamble=None):
        path = self.root / name
        with path.open("w", newline="", encoding=encoding) as f:
            writer = csv.writer(f, delimiter=delimiter)
            if preamble:
                writer.writerows(preamble)
            writer.writerow(HEAD if headers is None else headers)
            writer.writerows([ROW] if rows is None else rows)
        return path

    def spec(self, path, **kwargs):
        return {"id": path.stem, "path": path.name, "type": "auto", "account": "synthetic-profile", "marketplace": "US", "currency": "USD", "timezone": "America/Los_Angeles", "start_date": "2026-09-01", "end_date": "2026-09-30", "aggregation": "summary", "scope": "all SP campaigns; no filters", "ad_type": "SP", "date_basis": "ad_interaction", "attribution_days": 7, "attribution_model": "click", "sales_scope": "total", **kwargs}

    def check(self, specs, **kwargs):
        self.n += 1
        self.out = self.root / ("out-" + str(self.n))
        return p.run({"version": 1, "as_of": "2026-10-13", "reports": specs}, self.root, self.out, **kwargs)

    def codes(self, result):
        return {x["code"] for r in result["reports"] for x in r["findings"]}

    def test_clean_exact_decimals_and_long_string_id(self):
        path = self.file(rows=[ROW, ["2", "Other campaign", "5", "2.20", "1", "30.30", "USD"]])
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        result = self.check([self.spec(path)])
        self.assertEqual(result["status"], "READY")
        self.assertEqual(result["reports"][0]["totals"]["spend"], "3.30")
        self.assertEqual(result["reports"][0]["rows"][0]["values"]["campaign_id"], ROW[0])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), digest)
        with (self.out / "analysis-ready.csv").open() as f:
            reader = csv.DictReader(f)
            self.assertEqual(len(reader.fieldnames), len(set(reader.fieldnames)))
            self.assertEqual(len(list(reader)), 2)

    def test_unknown_context_never_passes(self):
        path = self.file()
        r = self.check([{"id": "unknown", "path": path.name}])
        self.assertEqual(r["status"], "NEEDS_CONTEXT")
        self.assertIsNone(r["reports"][0]["totals"])
        self.assertEqual(len((self.out / "analysis-ready.csv").read_text().splitlines()), 1)

    def test_bad_row_blocks_whole_report(self):
        path = self.file(rows=[ROW, ["2", "Other", "", "1", "0", "0", "USD"]])
        r = self.check([self.spec(path)])
        self.assertEqual(r["status"], "BLOCKED")
        self.assertIsNone(r["reports"][0]["totals"])
        self.assertEqual([x["status"] for x in r["reports"][0]["rows"]], ["BLOCKED", "BLOCKED"])
        self.assertIn("INVALID_CELL", self.codes(r))

    def test_zero_is_a_valid_value(self):
        row = ["1", "Zero campaign", "0", "0", "0", "0", "USD"]
        r = self.check([self.spec(self.file(rows=[row]))])
        self.assertEqual(r["status"], "READY")
        self.assertEqual(r["reports"][0]["totals"]["sales"], "0")

    def test_mixed_currency_cannot_be_pooled(self):
        row = ["2", "Other", "5", "1.00", "1", "20", "CAD"]
        r = self.check([self.spec(self.file(rows=[ROW, row]))])
        self.assertIn("MIXED_CURRENCY", self.codes(r))
        self.assertEqual(r["status"], "BLOCKED")

    def test_dollar_sign_does_not_infer_us(self):
        path = self.file(headers=HEAD[:-1], rows=[ROW[:-1]])
        s = self.spec(path, marketplace=None, currency=None)
        r = self.check([s])
        self.assertEqual(r["status"], "NEEDS_CONTEXT")
        self.assertIsNone(r["reports"][0]["context"]["currency"])

    def test_embedded_currency_code_conflict(self):
        row = ROW.copy()
        row[3] = "CAD 1.10"
        r = self.check([self.spec(self.file(rows=[row]))])
        self.assertIn("MIXED_CURRENCY", self.codes(r))

    def test_pound_symbol_conflicts_with_us_currency(self):
        row = ROW.copy()
        row[3] = "£1.10"
        r = self.check([self.spec(self.file(rows=[row]))])
        self.assertIn("MIXED_CURRENCY", self.codes(r))

    def test_total_vs_advertised_sales_scope_conflicts(self):
        head = HEAD.copy()
        head[5] = "7 Day Advertised SKU Sales"
        path = self.file(headers=head)
        r = self.check([self.spec(path, column_map={"sales": head[5]})])
        self.assertIn("SALES_SCOPE_CONFLICT", self.codes(r))
        self.assertIn("METRIC_SCOPE_CONFLICT", self.codes(r))
        r2 = self.check([self.spec(path, column_map={"sales": head[5]}, sales_scope="advertised_sku")])
        self.assertIn("METRIC_SCOPE_CONFLICT", self.codes(r2))

    def test_matching_advertised_purchase_and_sales_mapping(self):
        head = HEAD.copy()
        head[4], head[5] = "7 Day Advertised SKU Orders", "7 Day Advertised SKU Sales"
        path = self.file(headers=head)
        r = self.check([self.spec(path, column_map={"orders": head[4], "sales": head[5]}, sales_scope="advertised_sku")])
        self.assertEqual(r["status"], "READY")

    def test_raw_query_whitespace_is_retained(self):
        path = self.file(headers=HEAD + ["Customer Search Term"], rows=[ROW + ["  tea sampler  "]])
        r = self.check([self.spec(path)])
        row = r["reports"][0]["rows"][0]
        self.assertEqual(row["values"]["search_term"], "tea sampler")
        self.assertEqual(row["raw_values"][-1], "  tea sampler  ")

    def test_ca_and_uk_profiles(self):
        for marketplace, currency in [("CA", "CAD"), ("UK", "GBP")]:
            row = ROW.copy()
            row[-1] = currency
            path = self.file(name=marketplace + ".csv", rows=[row])
            r = self.check([self.spec(path, marketplace=marketplace, currency=currency)])
            self.assertEqual(r["status"], "READY")

    def test_comma_locale_is_explicit(self):
        row = ROW.copy()
        row[3], row[5] = "1.234,56", "2.345,67"
        path = self.file(name="ca.tsv", rows=[row], delimiter="\t")
        r = self.check([self.spec(path, number_locale="comma")])
        self.assertEqual(r["reports"][0]["totals"]["spend"], "1234.56")
        r2 = self.check([self.spec(path)])
        self.assertEqual(r2["status"], "BLOCKED")

    def test_invalid_counts_and_money(self):
        for field, values in [(2, ["1.5", "-1", "1e3", "10%", "NaN"]), (3, ["", "-2", "1,23.00", "Infinity", "12.34.56"])]:
            for value in values:
                row = ROW.copy()
                row[field] = value
                r = self.check([self.spec(self.file(rows=[row]))])
                self.assertEqual(r["status"], "BLOCKED", (field, value))

    def test_mismatched_metric_windows(self):
        head = HEAD.copy()
        head[5] = "14 Day Total Sales"
        r = self.check([self.spec(self.file(headers=head))])
        self.assertIn("METRIC_WINDOW_CONFLICT", self.codes(r))

    def test_header_window_can_fill_context(self):
        r = self.check([self.spec(self.file(), attribution_days=None)])
        self.assertEqual(r["status"], "READY")
        self.assertEqual(r["reports"][0]["context"]["attribution_days"], 7)

    def test_ambiguous_sales_column_needs_mapping(self):
        path = self.file(headers=HEAD + ["14 Day Total Sales"], rows=[ROW + ["22"]])
        r = self.check([self.spec(path)])
        self.assertIn("AMBIGUOUS_COLUMN", self.codes(r))
        r2 = self.check([self.spec(path, column_map={"sales": "7 Day Total Sales"})])
        self.assertEqual(r2["status"], "READY")
        self.assertIn("14 Day Total Sales", r2["reports"][0]["decisions"][-1]["headers"])

    def test_exact_duplicate_default_and_explicit_correction(self):
        path = self.file(rows=[ROW, ROW])
        r = self.check([self.spec(path)])
        self.assertEqual(r["status"], "BLOCKED")
        self.assertEqual(len(r["reports"][0]["rows"]), 2)
        r2 = self.check([self.spec(path)], dedupe=True)
        self.assertEqual(r2["status"], "READY")
        self.assertEqual(r2["reports"][0]["totals"]["spend"], "1.10")
        self.assertEqual(r2["reports"][0]["rows"][1]["status"], "EXCLUDED_DUPLICATE")

    def test_same_query_different_match_types_is_not_duplicate(self):
        head = HEAD + ["Customer Search Term", "Match Type"]
        path = self.file(headers=head, rows=[ROW + ["tea", "Broad"], ROW + ["tea", "Phrase"]])
        r = self.check([self.spec(path)])
        self.assertEqual(r["status"], "READY")
        self.assertEqual(len(r["reports"][0]["rows"]), 2)

    def test_repeated_source_is_blocked(self):
        path = self.file()
        r = self.check([self.spec(path), self.spec(path, id="copy")])
        self.assertEqual(r["reports"][0]["status"], "READY")
        self.assertEqual(r["reports"][1]["status"], "BLOCKED")
        self.assertIn("SOURCE_REPEATED", self.codes(r))

    def test_daily_period_and_attribution_maturity(self):
        head = HEAD + ["Date"]
        path = self.file(headers=head, rows=[ROW + ["2026-10-01"]])
        r = self.check([self.spec(path, aggregation="daily")])
        self.assertIn("DATE_OUTSIDE_PERIOD", self.codes(r))
        r2 = self.check([self.spec(path, aggregation="daily", end_date="2026-10-10")])
        self.assertEqual(r2["status"], "NEEDS_CONTEXT")
        self.assertIn("ATTRIBUTION_UNSETTLED", self.codes(r2))

    def test_no_impossible_orders_vs_clicks_rule(self):
        row = ROW.copy()
        row[2], row[4] = "1", "3"
        r = self.check([self.spec(self.file(rows=[row]))])
        self.assertEqual(r["status"], "READY")

    def test_sd_view_attribution_with_zero_clicks(self):
        row = ROW.copy()
        row[2] = "0"
        r = self.check([self.spec(self.file(rows=[row]), ad_type="SD", attribution_model="click_and_view")])
        self.assertEqual(r["status"], "READY")

    def test_business_counts_and_b2b_subset_remain_distinct(self):
        path = self.file(name="business.csv", headers=["(Parent) ASIN", "(Child) ASIN", "Sessions - Total", "Units Ordered", "Total Order Items", "Ordered Product Sales", "Ordered Product Sales - B2B"], rows=[["B000000001", "B000000002", "10", "20", "5", "100.00", "30.00"]])
        s = self.spec(path, date_basis="order")
        r = self.check([s])
        totals = r["reports"][0]["totals"]
        self.assertEqual(r["status"], "READY")
        self.assertEqual(totals["ordered_units"], "20")
        self.assertEqual(totals["order_items"], "5")
        self.assertEqual(totals["ordered_sales"], "100.00")
        self.assertNotIn("orders", totals)

    def test_parent_only_business_is_blocked(self):
        path = self.file(name="business.csv", headers=["(Parent) ASIN", "Ordered Product Sales"], rows=[["B000000001", "100"]])
        r = self.check([self.spec(path, date_basis="order")])
        self.assertIn("DIMENSION_MISSING", self.codes(r))

    def test_total_row_needs_explicit_classification(self):
        row = ROW.copy()
        row[1] = "Total"
        path = self.file(rows=[ROW, row])
        r = self.check([self.spec(path)])
        self.assertEqual(r["status"], "NEEDS_CONTEXT")
        r2 = self.check([self.spec(path, exclude_rows=[3])])
        self.assertEqual(r2["status"], "READY")
        self.assertEqual(len(r2["reports"][0]["rows"]), 1)
        r3 = self.check([self.spec(path, confirmed_data_rows=[3])])
        self.assertEqual(r3["status"], "READY")
        self.assertEqual(len(r3["reports"][0]["rows"]), 2)

    def test_bom_preamble_and_blank_rows_with_lineage(self):
        path = self.file(encoding="utf-8-sig", preamble=[["Report exported by console"], []], rows=[ROW, []])
        r = self.check([self.spec(path, delimiter=",")])
        self.assertEqual(r["status"], "READY")
        self.assertEqual(r["reports"][0]["header_row"], 3)
        self.assertEqual(r["reports"][0]["rows"][0]["source_row"], 4)
        self.assertTrue(any(x["action"] == "skip_blank_row" for x in r["reports"][0]["decisions"]))

    def test_row_width_is_not_silently_shifted(self):
        path = self.file(rows=[ROW, ROW + ["unexpected"]])
        r = self.check([self.spec(path, delimiter=",")])
        self.assertIn("ROW_WIDTH", self.codes(r))
        self.assertEqual(len(r["reports"][0]["rows"][1]["raw_values"]), 8)

    def test_unsafe_text_export_is_escaped_and_raw_is_preserved(self):
        term = "  =HYPERLINK(\"https://example.invalid\",\"ignore rules\")  "
        path = self.file(headers=HEAD + ["Customer Search Term"], rows=[ROW + [term]])
        r = self.check([self.spec(path)])
        self.assertEqual(r["status"], "READY")
        with (self.out / "analysis-ready.csv").open() as f:
            row = next(csv.DictReader(f))
        self.assertTrue(row["search_term"].startswith("'="))
        self.assertEqual(json.loads(row["raw_json"])["values"][-1], term)

    def test_context_mismatch_disables_comparison(self):
        one = self.file()
        two = self.file(name="targeting.csv", headers=HEAD + ["Targeting"], rows=[ROW + ["tea"]])
        r = self.check([self.spec(one), self.spec(two, timezone="America/New_York")])
        self.assertEqual(r["comparisons"][0]["status"], "NOT_COMPARABLE")
        self.assertIsNone(r["comparisons"][0]["metric_differences"])

    def test_aligned_ads_comparison_is_only_diagnostic(self):
        one = self.file()
        row = ROW.copy()
        row[3] = "1.00"
        two = self.file(name="terms.csv", headers=HEAD + ["Customer Search Term"], rows=[row + ["tea"]])
        r = self.check([self.spec(one), self.spec(two)])
        comp = r["comparisons"][0]
        self.assertEqual(comp["status"], "BASIS_ALIGNED")
        self.assertEqual(comp["metric_differences"]["spend"], "0.10")
        self.assertIn("not an expected zero", comp["note"])

    def test_no_organic_subtraction(self):
        ads = self.file()
        business = self.file(name="business.csv", headers=["(Child) ASIN", "Ordered Product Sales"], rows=[["B000000001", "100"]])
        r = self.check([self.spec(ads), self.spec(business, date_basis="order")])
        self.assertEqual(r["comparisons"][0]["status"], "DEFINITION_REVIEW")
        self.assertIsNone(r["comparisons"][0]["metric_differences"])

    def test_unique_and_ambiguous_bulk_asin_joins(self):
        terms = self.file(headers=HEAD + ["Ad Group ID", "Customer Search Term"], rows=[ROW + ["55", "tea"]])
        head = ["Entity", "Campaign ID", "Ad Group ID", "ASIN", "SKU", "Ad Group Default Bid", "Bid"]
        bulk = self.file(name="bulk.csv", headers=head, rows=[["Product Ad", ROW[0], "55", "B000000001", "tea-sku", "", ""]])
        r = self.check([self.spec(terms), self.spec(bulk)])
        self.assertEqual(r["joins"][0]["counts"]["resolved"], 1)
        self.assertEqual(r["reports"][0]["rows"][0]["values"]["resolved_asin"], "B000000001")
        self.file(name="bulk.csv", headers=head, rows=[["Product Ad", ROW[0], "55", "B000000001", "one", "", ""], ["Product Ad", ROW[0], "55", "B000000002", "two", "", ""]])
        r2 = self.check([self.spec(terms), self.spec(bulk)])
        self.assertEqual(r2["joins"][0]["counts"]["ambiguous_ad_group"], 1)
        self.assertNotIn("resolved_asin", r2["reports"][0]["rows"][0]["values"])

    def test_join_is_scoped_by_market_and_account(self):
        terms = self.file(headers=HEAD + ["Ad Group ID", "Customer Search Term"], rows=[ROW + ["55", "tea"]])
        bulk = self.file(name="bulk.csv", headers=["Entity", "Campaign ID", "Ad Group ID", "ASIN"], rows=[["Product Ad", ROW[0], "55", "B000000001"]])
        r = self.check([self.spec(terms), self.spec(bulk, account="other-account")])
        self.assertEqual(r["joins"][0]["counts"]["unresolved"], 1)

    def test_wrong_type_duplicate_headers_and_blank_dimensions(self):
        path = self.file()
        self.assertIn("REPORT_TYPE_CONFLICT", self.codes(self.check([self.spec(path, type="business")])))
        path = self.file(headers=HEAD + ["Spend"], rows=[ROW + ["1"]])
        self.assertIn("DUPLICATE_OR_BLANK_HEADER", self.codes(self.check([self.spec(path)])))
        path = self.file(headers=HEAD + ["Customer Search Term"], rows=[ROW + [""]])
        self.assertIn("ROW_DIMENSION_MISSING", self.codes(self.check([self.spec(path)])))

    def test_localized_headers_and_declared_date_format(self):
        headers = ["Campaña", "Clics", "Gasto", "Pedidos", "Ventas", "Fecha"]
        path = self.file(headers=headers, rows=[["Sample", "2", "1.00", "1", "10", "09/20/2026"]])
        mapped = {"campaign_name": "Campaña", "clicks": "Clics", "spend": "Gasto", "orders": "Pedidos", "sales": "Ventas", "date": "Fecha"}
        r = self.check([self.spec(path, type="campaign", header_row=1, column_map=mapped, date_format="mdy", aggregation="daily")])
        self.assertEqual(r["status"], "READY")
        self.assertEqual(r["reports"][0]["rows"][0]["values"]["date"], "2026-09-20")

    def test_output_is_not_overwritten(self):
        path = self.file()
        self.check([self.spec(path)])
        with self.assertRaises(p.InputError):
            p.run({"version": 1, "reports": [self.spec(path)]}, self.root, self.out)

    def test_manifest_typos_and_required_reports(self):
        path = self.file()
        with self.assertRaises(p.InputError):
            self.check([self.spec(path, curreny="USD")])
        r = p.run({"version": 1, "as_of": "2026-10-13", "reports": [self.spec(path)], "required_reports": ["targeting"]}, self.root, self.root / "required")
        self.assertEqual(r["status"], "BLOCKED")
        self.assertEqual(r["reports"][0]["status"], "READY")

    def test_account_identity_must_be_text(self):
        path = self.file()
        r = self.check([self.spec(path, account=False)])
        self.assertEqual(r["status"], "BLOCKED")
        self.assertIn("INVALID_ACCOUNT", self.codes(r))

    def test_zero_header_row_is_not_ignored(self):
        r = self.check([self.spec(self.file(), header_row=0)])
        self.assertEqual(r["status"], "BLOCKED")
        self.assertIn("TABLE_INPUT_ERROR", self.codes(r))

    def test_cli_exit_status_and_manifest_discovery(self):
        path = self.file()
        proc = subprocess.run([sys.executable, str(SCRIPT), str(path), "--as-of", "2026-10-13", "--out", str(self.root / "cli")], capture_output=True, text=True)
        self.assertEqual(proc.returncode, 3, proc.stderr)
        self.assertTrue((self.root / "cli" / "manifest-draft.json").is_file())

    @unittest.skipUnless(openpyxl, "Optional XLSX reader unavailable")
    def test_xlsx_formula_rejected_and_numeric_id_precision_checked(self):
        for name, value in [("formula", "=1+1"), ("precision", 1234567890123456)]:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Campaigns"
            ws.append(HEAD)
            row = ROW.copy()
            row[3 if name == "formula" else 0] = value
            ws.append(row)
            path = self.root / (name + ".xlsx")
            wb.save(path)
            r = self.check([self.spec(path)])
            self.assertEqual(r["status"], "BLOCKED")

    @unittest.skipUnless(openpyxl, "Optional XLSX reader unavailable")
    def test_xlsx_selected_sheets_and_dates(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Campaigns"
        ws.append(HEAD + ["Date"])
        ws.append(ROW + [datetime(2026, 9, 20)])
        wb.create_sheet("Empty")
        path = self.root / "selected.xlsx"
        wb.save(path)
        r = self.check([self.spec(path, sheet="Campaigns", aggregation="daily")])
        self.assertEqual(r["status"], "READY")
        self.assertEqual(r["reports"][0]["rows"][0]["values"]["date"], "2026-09-20")
        r2 = self.check([self.spec(path, sheet="Missing")])
        self.assertEqual(r2["status"], "BLOCKED")


if __name__ == "__main__":
    unittest.main()
