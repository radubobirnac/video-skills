"""Synthetic regressions for upload integrity and order-independent routing QA."""
import copy
import csv
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPT = Path(__file__).resolve().parents[1] / "bulk_qa.py"
spec = importlib.util.spec_from_file_location("bulk_qa", SCRIPT)
qa = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = qa
spec.loader.exec_module(qa)


def row(n, entity, **values):
    data = {"Product": "Sponsored Products", "Entity": entity}
    data.update({k.replace("_", " "): str(v) for k, v in values.items()})
    return qa.Row(n, data)


def snapshot():
    return [
        row(2, "Campaign", Campaign_ID="001", Campaign_Name="ACME | SP - DISC - Widget", State="enabled", Targeting_Type="Manual", Daily_Budget="20"),
        row(3, "Ad Group", Campaign_ID="001", Ad_Group_ID="010", Ad_Group_Name="Widget Exact", State="enabled"),
        row(4, "Product Ad", Campaign_ID="001", Ad_Group_ID="010", Ad_ID="500", ASIN="B000000001", State="enabled"),
        row(5, "Keyword", Campaign_ID="001", Ad_Group_ID="010", Keyword_ID="100", Keyword_Text="old widget", Match_Type="Exact", Bid="1", State="enabled"),
        row(6, "Campaign", Campaign_ID="002", Campaign_Name="ACME | SP - AUTO - Widget", State="enabled", Targeting_Type="Auto", Daily_Budget="10"),
        row(7, "Ad Group", Campaign_ID="002", Ad_Group_ID="020", Ad_Group_Name="Widget Auto", State="enabled"),
        row(8, "Product Ad", Campaign_ID="002", Ad_Group_ID="020", Ad_ID="501", ASIN="B000000001", State="enabled"),
        row(9, "Campaign", Campaign_ID="003", Campaign_Name="ACME | SP - DISC - Other", State="enabled", Targeting_Type="Manual", Daily_Budget="30"),
        row(10, "Ad Group", Campaign_ID="003", Ad_Group_ID="030", Ad_Group_Name="Other Exact", State="enabled"),
        row(11, "Product Ad", Campaign_ID="003", Ad_Group_ID="030", Ad_ID="502", ASIN="B000000002", State="enabled"),
    ]


def manifest():
    return {"account_id": "test-profile", "marketplace": "US", "currency": "USD", "snapshot_exported_at": "2026-10-14T08:00:00Z", "reviewed_at": "2026-10-14T10:00:00Z", "snapshot_complete": True}


def create(n=2, text="new widget", **values):
    base = dict(Operation="Create", Campaign_ID="001", Ad_Group_ID="010", Keyword_Text=text, Match_Type="Exact", Bid="1", State="enabled")
    base.update(values)
    return row(n, "Keyword", **base)


def negative(n=3, text="new widget", **values):
    base = dict(Operation="Create", Campaign_ID="002", Ad_Group_ID="020", Keyword_Text=text, Match_Type="Negative Exact", State="enabled")
    base.update(values)
    return row(n, "Negative Keyword", **base)


def update(**values):
    base = dict(Operation="Update", Campaign_ID="001", Ad_Group_ID="010", Keyword_ID="100", Bid="1.25")
    base.update(values)
    return row(2, "Keyword", **base)


class Checks(unittest.TestCase):
    def setUp(self):
        self.s, self.m = snapshot(), manifest()

    def check(self, proposals):
        return qa.analyze(self.s, proposals, self.m)

    def code(self, proposals, code, status="BLOCKED"):
        result = self.check(proposals)
        self.assertEqual(result["status"], status)
        self.assertIn(code, [i["code"] for i in result["global_issues"]] + [i["code"] for r in proposals for i in r.issues])

    def test_valid_exact_create(self):
        r = create()
        self.assertEqual(self.check([r])["status"], "READY")
        self.assertEqual(r.payload["Campaign ID"], "001")

    def test_parent_mismatch(self):
        self.code([create(Ad_Group_ID="020")], "AD_GROUP_ID")

    def test_unknown_campaign(self):
        self.code([create(Campaign_ID="999")], "CAMPAIGN_ID")

    def test_formatted_id(self):
        self.code([create(Campaign_ID="1.0")], "ID_FORMAT")

    def test_create_has_keyword_id(self):
        self.code([create(Keyword_ID="100")], "CREATE_ID")

    def test_auto_destination(self):
        self.code([create(Campaign_ID="002", Ad_Group_ID="020")], "TARGETING_TYPE")

    def test_paused_destination(self):
        self.s[1].data["State"] = "paused"
        self.code([create()], "DESTINATION_INACTIVE")

    def test_missing_product_mapping(self):
        self.s[2].data["ASIN"] = ""
        self.code([create()], "PRODUCT_MAPPING")

    def test_existing_exact_normalized(self):
        self.code([create(text=" OLD   WIDGET ")], "EXISTING_EXACT")

    def test_paused_exact_is_duplicate(self):
        self.s[3].data["State"] = "paused"
        self.code([create(text="old widget")], "EXISTING_EXACT")

    def test_archived_exact_warns(self):
        self.s[3].data["State"] = "archived"
        r = create(text="old widget")
        self.assertEqual(self.check([r])["status"], "READY")
        self.assertIn("ARCHIVED_EXACT", [i["code"] for i in r.issues])

    def test_duplicate_candidates_all_blocked(self):
        rows = [create(2), create(3, " NEW   WIDGET ")]
        self.code(rows, "DUPLICATE_NEW_EXACT")
        self.assertEqual([r.status for r in rows], ["BLOCKED", "BLOCKED"])

    def test_different_asin_can_use_same_term(self):
        self.assertEqual(self.check([create(text="old widget", Campaign_ID="003", Ad_Group_ID="030")])["status"], "READY")

    def test_same_asin_cross_campaign_duplicate(self):
        self.s[9].data["ASIN"] = "B000000001"
        self.code([create(text="old widget", Campaign_ID="003", Ad_Group_ID="030")], "EXISTING_EXACT")

    def test_twentyfive(self):
        self.assertEqual(self.check([create(n + 2, f"widget {n}") for n in range(25)])["status"], "READY")

    def test_twentysix_all_blocked(self):
        rows = [create(n + 2, f"widget {n}") for n in range(26)]
        self.code(rows, "EXACT_CAP")
        self.assertTrue(all(r.status == "BLOCKED" for r in rows))

    def test_match_type_policy(self):
        self.code([create(Match_Type="Broad")], "MATCH_TYPE")

    def test_negative_length(self):
        self.code([negative(text="one two three four five", Match_Type="Negative Phrase")], "KEYWORD_LENGTH")

    def test_negative_broad(self):
        self.code([negative(Match_Type="Negative Broad")], "NEGATIVE_MATCH_TYPE")

    def test_bad_amounts(self):
        for value in ["NaN", "Infinity", "0", "-1", "1,25"]:
            with self.subTest(value=value):
                self.code([create(Bid=value)], "BID")

    def test_bid_boundary(self):
        self.assertEqual(self.check([update()])["status"], "READY")
        self.code([update(Bid="1.2501")], "CHANGE_LIMIT")

    def test_bid_parent_mismatch(self):
        self.code([update(Campaign_ID="003", Ad_Group_ID="030")], "UPDATE_PARENT")

    def test_name_drift(self):
        self.code([update(Campaign_Name="renamed")], "NAME_DRIFT")

    def test_state_change(self):
        self.code([update(State="paused")], "UNEXPECTED_FIELD")

    def test_match_change(self):
        self.code([update(Match_Type="Broad")], "UNEXPECTED_FIELD")

    def test_unchanged_fields_are_omitted(self):
        r = update(State="enabled", Match_Type="Exact", Keyword_Text="old widget")
        self.assertEqual(self.check([r])["status"], "READY")
        self.assertNotIn("State", r.payload)
        self.assertNotIn("Match Type", r.payload)

    def test_unknown_nonempty_field(self):
        self.code([create(Surprise="value")], "UNEXPECTED_FIELD")

    def test_duplicate_update(self):
        second = update(Bid="1.10")
        second.number = 3
        self.code([update(), second], "DUPLICATE_UPDATE")

    def test_budget_boundary(self):
        r = row(2, "Campaign", Operation="Update", Campaign_ID="001", Daily_Budget="25")
        self.assertEqual(self.check([r])["status"], "READY")
        r.data["Daily Budget"] = "25.01"
        self.code([r], "CHANGE_LIMIT")

    def test_unsupported_archive(self):
        self.code([update(Operation="Archive")], "UNSUPPORTED_OPERATION")

    def test_unsupported_sb(self):
        r = create()
        r.data["Product"] = "Sponsored Brands"
        self.code([r], "UNSUPPORTED_OPERATION")

    def test_stale_snapshot(self):
        self.m["snapshot_exported_at"] = "2026-10-01T08:00:00Z"
        self.code([create()], "STALE_SNAPSHOT")

    def test_future_snapshot(self):
        self.m["snapshot_exported_at"] = "2026-10-15T08:00:00Z"
        self.code([create()], "STALE_SNAPSHOT")

    def test_naive_timestamp(self):
        self.m["snapshot_exported_at"] = "2026-10-14T08:00:00"
        self.code([create()], "TIMESTAMP", "NEEDS_CONTEXT")

    def test_incomplete_snapshot(self):
        self.m["snapshot_complete"] = False
        self.code([create()], "INCOMPLETE_SNAPSHOT", "NEEDS_CONTEXT")

    def test_missing_account(self):
        del self.m["account_id"]
        self.code([create()], "MISSING_CONTEXT", "NEEDS_CONTEXT")

    def test_uk_ca_currency(self):
        for marketplace, currency in [("UK", "GBP"), ("CA", "CAD")]:
            self.m.update(marketplace=marketplace, currency=currency)
            self.assertEqual(self.check([create()])["status"], "READY")
        self.m["currency"] = "USD"
        self.code([create()], "CURRENCY")

    def test_ambiguous_snapshot_id(self):
        r = copy.deepcopy(self.s[0])
        r.number = 12
        self.s.append(r)
        self.code([create()], "SNAPSHOT_DUPLICATE_ID")

    def test_missing_snapshot_state(self):
        self.s[3].data["State"] = ""
        self.code([create()], "SNAPSHOT_STATE")

    def test_existing_negative_blocks_new_exact(self):
        self.s.append(row(12, "Negative Keyword", Campaign_ID="001", Ad_Group_ID="010", Keyword_ID="300", Keyword_Text="new widget", Match_Type="Negative Exact", State="enabled"))
        self.code([create()], "NEGATIVE_CONFLICT")

    def test_phrase_negative_matches_whole_tokens(self):
        self.s.append(row(12, "Campaign Negative Keyword", Campaign_ID="001", Keyword_ID="300", Keyword_Text="widget", Match_Type="Negative Phrase", State="enabled"))
        self.code([create()], "NEGATIVE_CONFLICT")
        self.assertEqual(self.check([create(text="widgetry")])["status"], "READY")

    def test_proposed_negative_conflict_order_independent(self):
        for reverse in [False, True]:
            rows = [create(), negative(Campaign_ID="001", Ad_Group_ID="010")]
            if reverse:
                rows.reverse()
            self.code(rows, "NEGATIVE_CONFLICT")

    def test_protected_campaign(self):
        self.m["rules"] = {"protected_campaign_ids": ["002"]}
        self.code([negative(text="wasteful")], "PROTECTED_CAMPAIGN")

    def test_protected_term(self):
        self.m["rules"] = {"protected_terms": ["acme widget"]}
        self.code([negative(text="acme", Match_Type="Negative Phrase")], "PROTECTED_TERM")

    def test_plain_waste_negative(self):
        self.assertEqual(self.check([negative(text="irrelevant")])["status"], "READY")

    def test_harvest_link_required(self):
        self.code([create(), negative()], "HARVEST_LINK_REQUIRED", "NEEDS_CONTEXT")

    def test_phase1_defers_negative(self):
        self.m["harvest_links"] = [{"negative_row": 3, "positive_row": 2}]
        rows = [create(), negative()]
        self.assertEqual(self.check(rows)["status"], "PHASE1_READY")
        self.assertEqual([r.status for r in rows], ["READY", "DEFERRED"])

    def test_phase1_bad_destination(self):
        self.m["harvest_links"] = [{"negative_row": 3, "positive_row": 2}]
        self.code([create(Bid="-1"), negative()], "HARVEST_DESTINATION")

    def phase2(self):
        self.m["harvest_links"] = [{"negative_row": 3, "destination_keyword_id": "100", "delivery_evidence": {"checked_at": "2026-10-14T09:00:00Z", "window_start": "2026-10-14T00:00:00Z", "window_end": "2026-10-14T08:00:00Z", "keyword_id": "100", "impressions": 20, "budget_available": True, "eligible_asins": ["B000000001"], "reference": "Keyword report and console review"}}]
        return negative(text="old widget")

    def test_phase2_delivery(self):
        self.assertEqual(self.check([self.phase2()])["status"], "READY")

    def test_phase2_old_evidence(self):
        r = self.phase2()
        self.m["harvest_links"][0]["delivery_evidence"]["checked_at"] = "2026-10-10T09:00:00Z"
        self.code([r], "DELIVERY_EVIDENCE", "NEEDS_CONTEXT")

    def test_phase2_no_impressions(self):
        r = self.phase2()
        self.m["harvest_links"][0]["delivery_evidence"]["impressions"] = 0
        self.code([r], "DELIVERY_EVIDENCE", "NEEDS_CONTEXT")

    def test_phase2_no_budget(self):
        r = self.phase2()
        self.m["harvest_links"][0]["delivery_evidence"]["budget_available"] = False
        self.code([r], "DELIVERY_EVIDENCE", "NEEDS_CONTEXT")

    def test_phase2_wrong_asin(self):
        r = self.phase2()
        self.m["harvest_links"][0]["delivery_evidence"]["eligible_asins"] = ["B000000002"]
        self.code([r], "DELIVERY_EVIDENCE", "NEEDS_CONTEXT")

    def test_phase2_destination_negative(self):
        r = self.phase2()
        self.s.append(row(12, "Negative Keyword", Campaign_ID="001", Ad_Group_ID="010", Keyword_ID="300", Keyword_Text="old widget", Match_Type="Negative Exact", State="enabled"))
        self.code([r], "HARVEST_BLOCKED")

    def test_phase2_paused_destination(self):
        r = self.phase2()
        self.s[3].data["State"] = "paused"
        self.code([r], "HARVEST_INACTIVE")

    def test_phase2_wrong_term(self):
        r = self.phase2()
        r.data["Keyword Text"] = "different"
        self.code([r], "HARVEST_DESTINATION")

    def test_no_proposals(self):
        self.assertEqual(self.check([])["status"], "NO_CHANGES")

    def test_future_campaign_start(self):
        self.s[0].data["Start Date"] = "20261015"
        self.code([create()], "DESTINATION_INACTIVE")

    def test_expired_campaign(self):
        self.s[0].data["End Date"] = "2026-10-13"
        self.code([create()], "DESTINATION_INACTIVE")

    def test_phase1_multiple_source_products(self):
        self.s.append(row(12, "Product Ad", Campaign_ID="002", Ad_Group_ID="020", Ad_ID="503", ASIN="B000000003", State="enabled"))
        self.m["harvest_links"] = [{"negative_row": 3, "positive_row": 2}]
        self.code([create(), negative()], "HARVEST_PRODUCTS")

    def test_phase2_historical_impressions_not_current(self):
        r = self.phase2()
        self.m["harvest_links"][0]["delivery_evidence"]["window_start"] = "2026-09-01T00:00:00Z"
        self.code([r], "DELIVERY_EVIDENCE", "NEEDS_CONTEXT")

    def test_phase2_partial_asin_evidence(self):
        self.s.extend([row(12, "Product Ad", Campaign_ID="002", Ad_Group_ID="020", Ad_ID="503", ASIN="B000000003", State="enabled"), row(13, "Product Ad", Campaign_ID="001", Ad_Group_ID="010", Ad_ID="504", ASIN="B000000003", State="enabled")])
        self.code([self.phase2()], "DELIVERY_EVIDENCE", "NEEDS_CONTEXT")

    def test_phase2_fractional_impressions(self):
        r = self.phase2()
        self.m["harvest_links"][0]["delivery_evidence"]["impressions"] = "1.5"
        self.code([r], "DELIVERY_EVIDENCE", "NEEDS_CONTEXT")

    def test_archived_parent_update(self):
        self.s[0].data["State"] = "archived"
        self.code([update()], "ARCHIVED_PARENT")

    def test_invalid_harvest_row_type(self):
        self.m["harvest_links"] = [{"negative_row": [3], "positive_row": 2}]
        with self.assertRaises(qa.InputError):
            self.check([create(), negative()])

    def test_campaign_negative_harvest_scope(self):
        self.m["harvest_links"] = [{"negative_row": 3, "positive_row": 2}]
        neg = negative(Ad_Group_ID="")
        neg.data["Entity"] = "Campaign Negative Keyword"
        self.assertEqual(self.check([create(), neg])["status"], "PHASE1_READY")

    def test_row_account_conflict(self):
        self.s[0].data["Account ID"] = "another-profile"
        self.code([create()], "ROW_CONTEXT")


class Files(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.p = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def save(self, name, rows):
        headers = list(dict.fromkeys(k for r in rows for k in r.data))
        with (self.p / name).open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            writer.writerows(r.data for r in rows)

    def runfiles(self, proposals, m=None):
        self.save("snapshot.csv", snapshot())
        self.save("changes.csv", proposals)
        m = m or manifest()
        m.update(snapshot="snapshot.csv", changes="changes.csv")
        (self.p / "manifest.json").write_text(json.dumps(m))
        return qa.run(self.p / "manifest.json", self.p / "out")

    def test_roundtrip_and_source_hash(self):
        result = self.runfiles([create()])
        self.assertEqual(result["status"], "READY")
        import openpyxl
        wb = openpyxl.load_workbook(self.p / "out/checked-upload.xlsx")
        headers = [c.value for c in wb.active[1]]
        self.assertEqual(wb.active.cell(2, headers.index("Campaign ID") + 1).value, "001")
        self.assertTrue(all(c.data_type != "f" for rr in wb.active for c in rr))
        wb.close()
        self.assertEqual(result["inputs"]["snapshot"]["sha256"], qa.sha(self.p / "snapshot.csv"))

    def test_blocked_bundle_no_partial_upload(self):
        result = self.runfiles([create(), create(3, "other", Bid="0")])
        self.assertEqual(result["status"], "BLOCKED")
        self.assertFalse((self.p / "out/checked-upload.xlsx").exists())

    def test_phase1_workbook_excludes_source_negative(self):
        m = manifest()
        m["harvest_links"] = [{"negative_row": 3, "positive_row": 2}]
        result = self.runfiles([create(), negative()], m)
        self.assertEqual(result["upload"]["rows"], 1)
        self.assertEqual(result["status"], "PHASE1_READY")
        self.assertIn("new widget", (self.p / "out/deferred.csv").read_text())

    def test_no_overwrite(self):
        self.runfiles([create()])
        with self.assertRaises(qa.InputError):
            qa.run(self.p / "manifest.json", self.p / "out")

    def test_csv_safe_audit(self):
        self.assertEqual(qa.csv_safe(" =HYPERLINK(1)"), "' =HYPERLINK(1)")

    def test_xlsx_formula_rejected(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["Product", "Entity", "Campaign ID"])
        wb.active.append(["Sponsored Products", "Campaign", "=1+1"])
        wb.save(self.p / "input.xlsx")
        wb.close()
        with self.assertRaisesRegex(qa.InputError, "Formula"):
            qa.read_table(self.p / "input.xlsx")

    def test_xlsx_numeric_long_id_rejected(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["Product", "Entity", "Campaign ID"])
        wb.active.append(["Sponsored Products", "Campaign", 1234567890123456])
        wb.save(self.p / "input.xlsx")
        wb.close()
        with self.assertRaisesRegex(qa.InputError, "15 digits"):
            qa.read_table(self.p / "input.xlsx")

    def test_sheet_ambiguity_rejected(self):
        import openpyxl
        wb = openpyxl.Workbook()
        for ws in [wb.active, wb.create_sheet("second")]:
            ws.append(["Product", "Entity"])
            ws.append(["Sponsored Products", "Campaign"])
        wb.save(self.p / "input.xlsx")
        wb.close()
        with self.assertRaisesRegex(qa.InputError, "worksheet explicitly"):
            qa.read_table(self.p / "input.xlsx")

    def test_alias_collision_rejected(self):
        (self.p / "input.csv").write_text("Product,Entity,Campaign Name,campaign name\nSponsored Products,Campaign,a,b\n")
        with self.assertRaisesRegex(qa.InputError, "ambiguous"):
            qa.read_table(self.p / "input.csv")

    def amazon_layout(self):
        # Column layout of an Amazon SP bulk export: names on parent rows, informational copies on child rows.
        headers = ["Product", "Entity", "Operation", "Campaign ID", "Ad Group ID", "Portfolio ID", "Ad ID", "Keyword ID",
                   "Product Targeting ID", "Campaign Name", "Ad Group Name", "Campaign Name (Informational only)",
                   "Ad Group Name (Informational only)", "Portfolio Name (Informational only)", "Start Date", "End Date",
                   "Targeting Type", "State", "Campaign State (Informational only)", "Ad Group State (Informational only)",
                   "Daily Budget", "SKU", "ASIN (Informational only)", "Eligibility Status (Informational only)",
                   "Ad Group Default Bid", "Ad Group Default Bid (Informational only)", "Bid", "Keyword Text", "Match Type",
                   "Impressions", "Clicks", "Click-through Rate", "Spend", "Sales", "Orders", "Units", "Conversion Rate", "ACOS", "CPC", "ROAS"]
        def line(**values):
            return {h: values.get(h.replace(" ", "_").replace("(", "").replace(")", "").replace("-", "_"), "") for h in headers}
        info = {"Campaign_Name_Informational_only": "Widget Exact"}
        rows = [
            line(Product="Sponsored Products", Entity="Campaign", Campaign_ID="111", Campaign_Name="Widget Exact", Targeting_Type="Manual", State="enabled", Daily_Budget="20"),
            line(Product="Sponsored Products", Entity="Ad Group", Campaign_ID="111", Ad_Group_ID="222", Ad_Group_Name="Exact", State="enabled", Ad_Group_Default_Bid="1", **info),
            line(Product="Sponsored Products", Entity="Product Ad", Campaign_ID="111", Ad_Group_ID="222", Ad_ID="333", SKU="SKU-1", ASIN_Informational_only="B000000001", State="enabled", Ad_Group_Name_Informational_only="Exact", **info),
        ]
        with (self.p / "snapshot.csv").open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            writer.writerows(rows)

    def test_amazon_export_layout(self):
        self.amazon_layout()
        self.save("changes.csv", [row(2, "Keyword", Operation="Create", Campaign_ID="111", Ad_Group_ID="222", Keyword_Text="new widget", Match_Type="Exact", Bid="1", State="enabled", **{"Campaign Name (Informational only)": "Widget Exact"})])
        m = manifest()
        m.update(snapshot="snapshot.csv", changes="changes.csv")
        (self.p / "manifest.json").write_text(json.dumps(m))
        result = qa.run(self.p / "manifest.json", self.p / "out")
        self.assertEqual(result["status"], "READY", result["rows"])
        self.assertEqual(result["upload"]["rows"], 1)

    def test_informational_name_drift(self):
        self.amazon_layout()
        self.save("changes.csv", [row(2, "Keyword", Operation="Create", Campaign_ID="111", Ad_Group_ID="222", Keyword_Text="new widget", Match_Type="Exact", Bid="1", State="enabled", **{"Campaign Name (Informational only)": "Renamed"})])
        m = manifest()
        m.update(snapshot="snapshot.csv", changes="changes.csv")
        (self.p / "manifest.json").write_text(json.dumps(m))
        result = qa.run(self.p / "manifest.json", self.p / "out")
        self.assertEqual(result["status"], "BLOCKED")
        self.assertIn("NAME_DRIFT", [i["code"] for r in result["rows"] for i in r["issues"]])

    def test_unknown_rule_rejected(self):
        m = manifest()
        m["rules"] = {"max_new_exact": 999}
        with self.assertRaises(qa.InputError):
            qa.analyze(snapshot(), [create()], m)


if __name__ == "__main__":
    unittest.main()
