#!/usr/bin/env python3
# Copyright (c) 2026 Radu Bobirnac / Profit First PPC · https://profitfirstppc.com
"""Read-only Sponsored Products bulk QA; emit a separate, reviewed upload workbook."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

VERSION = "1.0.0"
ATTRIBUTION = {
    "creator": "Radu Bobirnac", "publisher": "Profit First PPC",
    "skill": "amazon-ppc-bulk-qa", "website": "https://profitfirstppc.com",
    "referral_url": "https://profitfirstppc.com/tools/ai-skills/?utm_source=amazon-ppc-bulk-qa&utm_medium=referral&utm_campaign=amazon_skills&utm_content=report#amazon-ppc-bulk-qa",
}
CREDIT = "Prepared with PPC Bulk File QA by **Radu Bobirnac · [Profit First PPC](" + ATTRIBUTION["referral_url"] + ")**."
CANON = ["Product", "Entity", "Operation", "Campaign ID", "Ad Group ID",
         "Portfolio ID", "Ad ID", "Keyword ID", "Product Targeting ID",
         "Campaign Name", "Ad Group Name", "Start Date", "End Date",
         "Targeting Type", "State", "Daily Budget", "SKU", "ASIN",
         "Ad Group Default Bid", "Bid", "Keyword Text", "Match Type",
         "Bidding Strategy", "Placement", "Percentage", "Product Targeting Expression"]
READONLY = {"Impressions", "Clicks", "Spend", "Sales", "Orders", "Units",
            "ACOS", "ROAS", "CPC", "CTR", "Conversion Rate"}
IDS = {x for x in CANON if x.endswith(" ID")}
ALIASES = {x.casefold(): x for x in CANON + sorted(READONLY)}
ALIASES.update({x.casefold(): x for x in ["Account ID", "Profile ID", "Marketplace", "Currency", "Campaign State", "Ad Group State"]})
for _name in ["Campaign Name", "Ad Group Name", "Campaign State", "Ad Group State"]:
    ALIASES[(_name + " (Informational only)").casefold()] = _name
DEFAULTS = {"new_exact_only": True, "max_new_exacts_per_campaign": 25,
            "snapshot_max_age_hours": 48, "delivery_max_age_hours": 24,
            "max_bid_change_pct": 25, "max_budget_change_pct": 25,
            "duplicate_scope": "same_asin", "protected_terms": [],
            "protected_campaign_ids": []}
UPLOAD = ["Product", "Entity", "Operation", "Campaign ID", "Ad Group ID",
          "Keyword ID", "State", "Daily Budget", "Bid", "Keyword Text", "Match Type"]


def norm(value):
    return " ".join(unicodedata.normalize("NFC", str(value)).casefold().split())


def amount(value):
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("Use an unformatted decimal number with a dot.") from exc
    if not result.is_finite() or result <= 0:
        raise ValueError("Value must be finite and greater than zero.")
    return result


def instant(value):
    result = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if result.tzinfo is None:
        raise ValueError("Timestamp must include a UTC offset.")
    return result.astimezone(timezone.utc)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class InputError(ValueError):
    pass


@dataclass
class Row:
    number: int
    data: dict
    sheet: str = ""
    issues: list = field(default_factory=list)
    before: dict = field(default_factory=dict)
    payload: dict = field(default_factory=dict)

    def get(self, key):
        return self.data.get(key, "")

    def add(self, code, message, severity="BLOCKED"):
        item = {"code": code, "severity": severity, "message": message}
        if item not in self.issues:
            self.issues.append(item)

    @property
    def status(self):
        for status in ["BLOCKED", "NEEDS_CONTEXT", "DEFERRED"]:
            if any(i["severity"] == status for i in self.issues):
                return status
        return "READY"


def read_table(path, sheet=None):
    """Keep IDs and names as text, refuse formulas and ambiguous schemas."""
    path = Path(path)
    notes = []
    if path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise InputError("Install openpyxl to read or export XLSX files.") from exc
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet:
                if sheet not in wb.sheetnames:
                    raise InputError(f"Worksheet {sheet!r} not found.")
                ws = wb[sheet]
            else:
                candidates = []
                for candidate in wb:
                    first = next(candidate.iter_rows(), ())
                    headers = [str(c.value or "").strip().casefold() for c in first]
                    if "entity" not in headers or "product" not in headers:
                        continue
                    pi = headers.index("product")
                    products = {norm(row[pi].value or "") for row in candidate.iter_rows(min_row=2)
                                if len(row) > pi and row[pi].value}
                    if "sponsored products" in products or not products:
                        candidates.append(candidate)
                if len(candidates) != 1:
                    raise InputError("Select a worksheet explicitly: zero or multiple SP entity sheets found.")
                ws = candidates[0]
                notes = [f"Ignored worksheet: {x}" for x in wb.sheetnames if x != ws.title]
            values = []
            for rr in ws.iter_rows():
                vals = []
                for cell in rr:
                    if cell.data_type == "f":
                        raise InputError(f"Formula at {ws.title}!{cell.coordinate}; provide a values-only export.")
                    val = cell.value
                    if isinstance(val, (int, float)) and len(str(int(abs(val)))) > 15:
                        raise InputError(f"Numeric cell {ws.title}!{cell.coordinate} exceeds 15 digits; recover IDs from a text export.")
                    if isinstance(val, datetime):
                        val = val.isoformat()
                    elif isinstance(val, float) and val.is_integer():
                        val = int(val)
                    vals.append("" if val is None else str(val))
                values.append(vals)
            sheet = ws.title
        finally:
            wb.close()
    elif path.suffix.lower() in {".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            values = list(csv.reader(handle, delimiter="\t" if path.suffix.lower() == ".tsv" else ","))
        sheet = ""
    else:
        raise InputError("Use a CSV, TSV or XLSX file; macro workbooks are unsupported.")
    if not values:
        raise InputError("Input is empty.")
    notes.append("Original headers: " + json.dumps(values[0], ensure_ascii=False))
    headers = [ALIASES.get(x.strip().casefold(), x.strip()) for x in values[0]]
    nonempty = [h for h in headers if h]
    if len(nonempty) != len(set(nonempty)):
        raise InputError("Duplicate or ambiguous column names after alias normalization.")
    if not {"Product", "Entity"}.issubset(headers):
        raise InputError("Missing Product/Entity columns; provide a bulk configuration export.")
    rows = []
    for number, vals in enumerate(values[1:], 2):
        if not any(v for v in vals):
            continue
        if len(vals) > len(headers) and any(vals[len(headers):]):
            raise InputError(f"Row {number} has cells beyond the header.")
        vals += [""] * (len(headers) - len(vals))
        if any(v for h, v in zip(headers, vals) if not h):
            raise InputError(f"Row {number} has data in an unnamed column.")
        rows.append(Row(number, {h: v for h, v in zip(headers, vals) if h}, sheet))
    return rows, headers, notes


def issue(code, message, severity="BLOCKED"):
    return {"code": code, "message": message, "severity": severity}


def validate_rules(manifest):
    supplied = manifest.get("rules", {})
    if not isinstance(supplied, dict) or set(supplied) - (set(DEFAULTS) | {"max_new_bid", "max_daily_budget"}):
        raise InputError("Unknown rules or invalid rules object; see manifest reference.")
    rules = {**DEFAULTS, **supplied}
    if not isinstance(rules["new_exact_only"], bool):
        raise InputError("new_exact_only must be a boolean.")
    cap = rules["max_new_exacts_per_campaign"]
    if isinstance(cap, bool) or not isinstance(cap, int) or cap < 0:
        raise InputError("max_new_exacts_per_campaign must be a nonnegative integer.")
    for k in ["snapshot_max_age_hours", "delivery_max_age_hours", "max_bid_change_pct", "max_budget_change_pct", "max_new_bid", "max_daily_budget"]:
        if k in rules:
            amount(rules[k])
    if rules["duplicate_scope"] not in {"ad_group", "campaign", "same_asin"}:
        raise InputError("duplicate_scope must be ad_group, campaign or same_asin.")
    for k in ["protected_terms", "protected_campaign_ids"]:
        if not isinstance(rules[k], list) or any(not isinstance(x, str) or not x for x in rules[k]):
            raise InputError(f"{k} must be a list of nonempty strings.")
    return rules


def check_context(manifest, rules):
    errors = []
    required = ["account_id", "marketplace", "currency", "snapshot_exported_at", "snapshot_complete"]
    for key in required:
        if key not in manifest or manifest[key] in (None, ""):
            errors.append(issue("MISSING_CONTEXT", f"Supply {key} from verified export context.", "NEEDS_CONTEXT"))
    market = manifest.get("marketplace")
    if manifest.get("account_id") and (not isinstance(manifest["account_id"], str) or not manifest["account_id"].strip()):
        errors.append(issue("ACCOUNT_ID", "account_id must be verified nonempty text."))
    if market and market not in {"US", "UK", "CA"}:
        errors.append(issue("MARKETPLACE", "Only US, UK and CA are supported."))
    if market in {"US", "UK", "CA"} and manifest.get("currency") != {"US": "USD", "UK": "GBP", "CA": "CAD"}[market]:
        errors.append(issue("CURRENCY", "Currency does not match the marketplace."))
    if "snapshot_complete" in manifest and manifest["snapshot_complete"] is not True:
        errors.append(issue("INCOMPLETE_SNAPSHOT", "Confirm a complete configuration export, including zero-impression and paused entities.", "NEEDS_CONTEXT"))
    try:
        now = instant(manifest.get("reviewed_at", datetime.now(timezone.utc).isoformat()))
        if manifest.get("snapshot_exported_at"):
            age = (now - instant(manifest["snapshot_exported_at"])).total_seconds() / 3600
            if age < 0 or Decimal(str(age)) > amount(rules["snapshot_max_age_hours"]):
                errors.append(issue("STALE_SNAPSHOT", "Snapshot is future-dated or exceeds the configured age limit."))
    except (ValueError, TypeError):
        now = datetime.now(timezone.utc)
        errors.append(issue("TIMESTAMP", "Use valid ISO timestamps with timezone offsets.", "NEEDS_CONTEXT"))
    return errors, now


def term_check(row, negative=False):
    text = row.get("Keyword Text")
    mt = norm(row.get("Match Type"))
    words = len(text.split())
    maximum = 4 if negative and mt == "negative phrase" else 10
    if not text.strip() or len(text) > 80 or words > maximum:
        row.add("KEYWORD_LENGTH", f"Keyword needs 1–{maximum} words and at most 80 characters.")
    if any(c in text for c in ['?', '"', '&', '\\']) or any(ord(c) < 32 for c in text):
        row.add("KEYWORD_CHARACTERS", "Keyword contains documented unsupported punctuation or control characters.")


def blocked_by(term, negative):
    a, b = norm(term), norm(negative.get("Keyword Text"))
    mt = norm(negative.get("Match Type"))
    if mt == "negative exact":
        return a == b
    return mt == "negative phrase" and bool(b) and f" {b} " in f" {a} "


def same_value(key, a, b):
    if key in {"State", "Match Type", "Targeting Type", "Product", "Entity"}:
        return norm(a) == norm(b)
    if key in {"Bid", "Daily Budget", "Ad Group Default Bid"}:
        try:
            return Decimal(a) == Decimal(b)
        except InvalidOperation:
            return a == b
    return a == b


def analyze(snapshot, proposals, manifest, headers=None):
    rules = validate_rules(manifest)
    global_issues, now = check_context(manifest, rules)
    maps = {"Campaign": {}, "Ad Group": {}, "Keyword": {}}
    idfields = {"Campaign": "Campaign ID", "Ad Group": "Ad Group ID", "Keyword": "Keyword ID"}
    negatives, ads, positives = [], [], []
    for row in snapshot:
        if norm(row.get("Product")) != "sponsored products":
            continue
        entity = row.get("Entity").strip().title()
        if entity in maps:
            key = row.get(idfields[entity])
            if not key or not re.fullmatch(r"[0-9]+", key):
                global_issues.append(issue("SNAPSHOT_ID", f"Snapshot row {row.number}: missing or invalid {idfields[entity]}."))
            elif key in maps[entity]:
                global_issues.append(issue("SNAPSHOT_DUPLICATE_ID", f"Repeated {entity} ID {key}; resolve snapshot ambiguity."))
            else:
                maps[entity][key] = row
        if entity == "Keyword":
            positives.append(row)
        elif entity in {"Negative Keyword", "Campaign Negative Keyword"}:
            negatives.append(row)
        elif entity == "Product Ad":
            ads.append(row)
    campaigns, groups, keywords = maps["Campaign"], maps["Ad Group"], maps["Keyword"]
    if not campaigns or not groups:
        global_issues.append(issue("SNAPSHOT_STRUCTURE", "Snapshot must contain campaigns and ad groups."))
    for row in snapshot:
        if norm(row.get("Product")) != "sponsored products":
            continue
        entity = norm(row.get("Entity"))
        cid, gid = row.get("Campaign ID"), row.get("Ad Group ID")
        if entity in {"campaign", "ad group", "keyword", "product ad", "negative keyword", "campaign negative keyword"} and norm(row.get("State")) not in {"enabled", "paused", "archived"}:
            global_issues.append(issue("SNAPSHOT_STATE", f"Snapshot row {row.number}: missing or unknown State."))
        if entity in {"keyword", "product ad", "negative keyword"} and not gid:
            global_issues.append(issue("SNAPSHOT_PARENT", f"Snapshot row {row.number}: missing Ad Group ID."))
        if entity in {"keyword", "negative keyword", "campaign negative keyword"} and (not row.get("Keyword Text") or norm(row.get("Match Type")) not in ({"exact", "phrase", "broad"} if entity == "keyword" else {"negative exact", "negative phrase"})):
            global_issues.append(issue("SNAPSHOT_KEYWORD", f"Snapshot row {row.number}: incomplete keyword definition."))
        if entity != "campaign" and cid not in campaigns:
            global_issues.append(issue("SNAPSHOT_PARENT", f"Snapshot row {row.number}: unknown parent campaign."))
        if gid and entity != "ad group" and (gid not in groups or groups[gid].get("Campaign ID") != cid):
            global_issues.append(issue("SNAPSHOT_PARENT", f"Snapshot row {row.number}: invalid ad group parent."))
        for column in ["Marketplace", "Currency", "Account ID", "Profile ID"]:
            expected = manifest.get({"Marketplace": "marketplace", "Currency": "currency", "Account ID": "account_id", "Profile ID": "account_id"}[column])
            if row.get(column) and row.get(column) != expected:
                global_issues.append(issue("ROW_CONTEXT", f"Snapshot row {row.number}: {column} conflicts with manifest."))
    products = defaultdict(set)
    enabled_products = defaultdict(set)
    for ad in ads:
        if norm(ad.get("State")) != "archived" and ad.get("ASIN"):
            products[ad.get("Ad Group ID")].add(ad.get("ASIN").upper())
        if norm(ad.get("State")) == "enabled" and ad.get("ASIN"):
            enabled_products[ad.get("Ad Group ID")].add(ad.get("ASIN").upper())

    def active_destination(row):
        cid, gid = row.get("Campaign ID"), row.get("Ad Group ID")
        camp, group = campaigns.get(cid), groups.get(gid)
        if not camp or not group:
            return False
        try:
            budget_ok = amount(camp.get("Daily Budget")) > 0
        except ValueError:
            budget_ok = False
        dates_ok = True
        for key in ["Start Date", "End Date"]:
            value = camp.get(key)
            if value:
                try:
                    date = datetime.strptime(value, "%Y%m%d").date() if re.fullmatch(r"\d{8}", value) else datetime.fromisoformat(value.replace("Z", "+00:00")).date()
                    dates_ok &= date <= now.date() if key == "Start Date" else date >= now.date()
                except ValueError:
                    dates_ok = False
        return norm(camp.get("State")) == norm(group.get("State")) == "enabled" and budget_ok and dates_ok and bool(enabled_products[gid])

    def overlaps(a, b):
        scope = rules["duplicate_scope"]
        if scope == "ad_group":
            return a.get("Ad Group ID") == b.get("Ad Group ID")
        if scope == "campaign":
            return a.get("Campaign ID") == b.get("Campaign ID")
        return a.get("Ad Group ID") == b.get("Ad Group ID") or bool(products[a.get("Ad Group ID")] & products[b.get("Ad Group ID")])

    def negative_scope(neg, pos):
        return neg.get("Campaign ID") == pos.get("Campaign ID") and (norm(neg.get("Entity")) == "campaign negative keyword" or neg.get("Ad Group ID") == pos.get("Ad Group ID"))

    creates, newneg = [], []
    seen_targets = defaultdict(list)
    for row in proposals:
        row.issues.clear()
        row.before.clear()
        row.payload.clear()
        entity = row.get("Entity").strip().title()
        operation = row.get("Operation").strip().title()
        cid, gid = row.get("Campaign ID"), row.get("Ad Group ID")
        camp, group = campaigns.get(cid), groups.get(gid)
        if norm(row.get("Product")) != "sponsored products" or (entity, operation) not in {
            ("Keyword", "Create"), ("Keyword", "Update"), ("Campaign", "Update"),
            ("Negative Keyword", "Create"), ("Campaign Negative Keyword", "Create")
        }:
            row.add("UNSUPPORTED_OPERATION", "V1 supports SP keyword creates/bid updates, budget updates and negative keyword creates.")
            continue
        for k in IDS:
            if row.get(k) and not re.fullmatch(r"[0-9]+", row.get(k)):
                row.add("ID_FORMAT", f"{k} must be the original numeric text ID, without formatting.")
        if not camp:
            row.add("CAMPAIGN_ID", "Campaign ID is absent from the snapshot.")
        if entity in {"Keyword", "Negative Keyword"} and (not group or group.get("Campaign ID") != cid):
            row.add("AD_GROUP_ID", "Ad Group ID is missing or does not belong to this campaign.")
        if entity in {"Campaign", "Campaign Negative Keyword"} and gid:
            row.add("CAMPAIGN_SCOPE", "Campaign-level rows must have an empty Ad Group ID.")
        if any(parent and norm(parent.get("State")) == "archived" for parent in [camp, group]):
            row.add("ARCHIVED_PARENT", "These operations cannot target an archived campaign/ad group.")
        for key, baseline in [("Campaign Name", camp), ("Ad Group Name", group)]:
            if row.get(key) and (baseline is None or row.get(key) != baseline.get(key)):
                row.add("NAME_DRIFT", f"{key} differs from the snapshot; names are preserved exactly.")
        row.payload = {"Product": "Sponsored Products", "Entity": entity, "Operation": operation, "Campaign ID": cid}
        if gid:
            row.payload["Ad Group ID"] = gid
        allowed = {"Product", "Entity", "Operation", "Campaign ID", "Ad Group ID", "Campaign Name", "Ad Group Name"}
        baseline = None
        if operation == "Create":
            if row.get("Keyword ID"):
                row.add("CREATE_ID", "New keywords and negatives must have an empty Keyword ID.")
            allowed |= {"Keyword ID", "Keyword Text", "Match Type", "State"}
            if norm(row.get("State")) != "enabled":
                row.add("CREATE_STATE", "New rows must explicitly use State enabled.")
            row.payload["State"] = "enabled"
            row.payload["Keyword Text"] = row.get("Keyword Text")
            term_check(row, entity != "Keyword")
            mt = norm(row.get("Match Type"))
            if entity == "Keyword":
                creates.append(row)
                allowed.add("Bid")
                accepted = {"exact"} if rules["new_exact_only"] else {"exact", "phrase", "broad"}
                if mt not in accepted:
                    row.add("MATCH_TYPE", "New keyword match type is outside configured policy.")
                row.payload["Match Type"] = mt.title()
                try:
                    bid = amount(row.get("Bid"))
                    row.payload["Bid"] = str(bid)
                    if "max_new_bid" in rules and bid > amount(rules["max_new_bid"]):
                        row.add("NEW_BID_LIMIT", "Bid exceeds the configured new-keyword ceiling.")
                except ValueError as exc:
                    row.add("BID", str(exc))
                if not camp or norm(camp.get("Targeting Type")) != "manual":
                    row.add("TARGETING_TYPE", "Destination must be an existing manual-targeting campaign.")
                if not active_destination(row):
                    row.add("DESTINATION_INACTIVE", "Destination needs enabled campaign/ad group/product ads and a positive daily budget.")
                if rules["duplicate_scope"] == "same_asin" and (not products[gid] or any(not products[x.get("Ad Group ID")] for x in positives if norm(x.get("State")) != "archived" and norm(x.get("Match Type")) == "exact")):
                    row.add("PRODUCT_MAPPING", "ASIN coverage is incomplete; same-ASIN duplicate checks cannot be certified.", "NEEDS_CONTEXT")
            else:
                newneg.append(row)
                if mt not in {"negative exact", "negative phrase"}:
                    row.add("NEGATIVE_MATCH_TYPE", "Use Negative Exact or Negative Phrase.")
                row.payload["Match Type"] = mt.title()
                if cid in rules["protected_campaign_ids"]:
                    row.add("PROTECTED_CAMPAIGN", "Negative targets a campaign protected by your policy.")
                if any(blocked_by(x, row) for x in rules["protected_terms"]):
                    row.add("PROTECTED_TERM", "Negative would block a protected term.")
        else:
            key = row.get("Keyword ID") if entity == "Keyword" else cid
            seen_targets[(entity, key)].append(row)
            baseline = keywords.get(key) if entity == "Keyword" else camp
            if not baseline:
                row.add("UPDATE_ID", "Update ID is absent from the snapshot.")
            elif entity == "Keyword" and (baseline.get("Campaign ID") != cid or baseline.get("Ad Group ID") != gid):
                row.add("UPDATE_PARENT", "Keyword ID belongs to a different campaign/ad group.")
            else:
                row.before = dict(baseline.data)
                if norm(baseline.get("State")) == "archived":
                    row.add("ARCHIVED_TARGET", "Archived entities cannot receive these updates.")
            if entity == "Keyword":
                allowed.add("Keyword ID")
                row.payload["Keyword ID"] = key
            metric = "Bid" if entity == "Keyword" else "Daily Budget"
            allowed.add(metric)
            try:
                new = amount(row.get(metric))
                old = amount(baseline.get(metric)) if baseline else None
                row.payload[metric] = str(new)
                if old is not None:
                    pct = abs(new - old) * 100 / old
                    limit = amount(rules["max_bid_change_pct" if entity == "Keyword" else "max_budget_change_pct"])
                    if pct > limit:
                        row.add("CHANGE_LIMIT", f"{metric} change {pct:.2f}% exceeds configured {limit}%.")
                    if new == old:
                        row.add("NO_CHANGE", "Update has no actual change.")
                if metric == "Daily Budget" and "max_daily_budget" in rules and new > amount(rules["max_daily_budget"]):
                    row.add("BUDGET_LIMIT", "Budget exceeds the configured ceiling.")
            except ValueError as exc:
                row.add("AMOUNT", str(exc))
        for column, value in row.data.items():
            if not value or column in allowed:
                continue
            if baseline and column in baseline.data and same_value(column, value, baseline.get(column)):
                row.add("OMITTED_UNCHANGED", f"Omitted unchanged field {column} from upload.", "INFO")
            else:
                row.add("UNEXPECTED_FIELD", f"Nonempty field {column} is outside the permitted operation or differs from baseline.")
    for target, rows in seen_targets.items():
        if len(rows) > 1:
            for row in rows:
                row.add("DUPLICATE_UPDATE", f"Multiple updates to {target}; consolidate into one proposal.")

    exacts = [x for x in creates if norm(x.get("Match Type")) == "exact"]
    for row in exacts:
        for existing in positives:
            if norm(existing.get("Match Type")) == "exact" and norm(existing.get("Keyword Text")) == norm(row.get("Keyword Text")) and overlaps(row, existing):
                row.add("EXISTING_EXACT" if norm(existing.get("State")) != "archived" else "ARCHIVED_EXACT",
                        f"Matching snapshot exact keyword {existing.get('Keyword ID')} in duplicate scope.",
                        "BLOCKED" if norm(existing.get("State")) != "archived" else "WARNING")
        for other in exacts:
            if other is not row and norm(other.get("Keyword Text")) == norm(row.get("Keyword Text")) and overlaps(row, other):
                row.add("DUPLICATE_NEW_EXACT", f"Proposal row {other.number} duplicates this exact in configured scope.")
    counts = Counter(x.get("Campaign ID") for x in exacts)
    for row in exacts:
        if counts[row.get("Campaign ID")] > rules["max_new_exacts_per_campaign"]:
            row.add("EXACT_CAP", "Campaign exceeds the configured number of new exacts; no arbitrary subset is retained.")
    for neg in newneg:
        for existing in negatives:
            if norm(existing.get("State")) != "archived" and norm(existing.get("Entity")) == norm(neg.get("Entity")) and existing.get("Campaign ID") == neg.get("Campaign ID") and existing.get("Ad Group ID") == neg.get("Ad Group ID") and norm(existing.get("Match Type")) == norm(neg.get("Match Type")) and norm(existing.get("Keyword Text")) == norm(neg.get("Keyword Text")):
                neg.add("EXISTING_NEGATIVE", "Negative already exists at this scope.")
        for other in newneg:
            if other is not neg and all(norm(other.get(k)) == norm(neg.get(k)) for k in ["Entity", "Campaign ID", "Ad Group ID", "Match Type", "Keyword Text"]):
                neg.add("DUPLICATE_NEGATIVE", f"Duplicate negative at proposal row {other.number}.")
    livepositives = [p for p in positives if norm(p.get("State")) == "enabled"] + creates
    for neg in negatives + newneg:
        if neg in negatives and norm(neg.get("State")) != "enabled":
            continue
        for pos in livepositives:
            if negative_scope(neg, pos) and blocked_by(pos.get("Keyword Text"), neg):
                exact = norm(pos.get("Match Type")) == "exact"
                if pos in creates:
                    pos.add("NEGATIVE_CONFLICT", f"Exact/text blocked by negative at {'proposal' if neg in newneg else 'snapshot'} row {neg.number}.")
                if neg in newneg:
                    neg.add("NEGATIVE_CONFLICT" if exact else "NEGATIVE_OVERLAP",
                            f"Negative overlaps {'exact' if exact else 'phrase/broad'} keyword {pos.get('Keyword ID') or 'proposal row ' + str(pos.number)}.",
                            "BLOCKED" if exact else "WARNING")

    links = manifest.get("harvest_links", [])
    if not isinstance(links, list):
        raise InputError("harvest_links must be an array.")
    bynumber = {r.number: r for r in proposals}
    linked = set()
    for link in links:
        if not isinstance(link, dict):
            raise InputError("Each harvest link must be an object.")
        for selector in ["negative_row", "positive_row"]:
            if selector in link and (type(link[selector]) is not int or link[selector] < 2):
                raise InputError(f"{selector} must be a physical row number of at least 2.")
        neg = bynumber.get(link.get("negative_row"))
        if not neg or neg not in newneg or norm(neg.get("Match Type")) != "negative exact" or neg.number in linked:
            global_issues.append(issue("HARVEST_LINK", "Harvest link needs one unique Negative Exact proposal row."))
            continue
        linked.add(neg.number)
        if ("positive_row" in link) == ("destination_keyword_id" in link):
            neg.add("HARVEST_LINK", "Supply exactly one positive_row or destination_keyword_id.")
            continue
        dest = bynumber.get(link.get("positive_row")) if "positive_row" in link else keywords.get(str(link.get("destination_keyword_id")))
        if not dest or ("positive_row" in link and dest not in exacts) or norm(dest.get("Match Type")) != "exact" or norm(dest.get("Keyword Text")) != norm(neg.get("Keyword Text")):
            neg.add("HARVEST_DESTINATION", "Destination must be the corresponding exact keyword.")
            continue
        sourceproducts = products[neg.get("Ad Group ID")]
        if norm(neg.get("Entity")) == "campaign negative keyword":
            sourceproducts = set().union(*(v for g, v in products.items() if groups.get(g) and groups[g].get("Campaign ID") == neg.get("Campaign ID")))
        if not sourceproducts or not sourceproducts <= enabled_products[dest.get("Ad Group ID")]:
            neg.add("HARVEST_PRODUCTS", "Destination must cover every ASIN affected by the source negative with enabled product ads.")
        if negative_scope(neg, dest):
            neg.add("HARVEST_SCOPE", "Source negative would block its own destination.")
        if "positive_row" in link:
            if dest.status != "READY":
                neg.add("HARVEST_DESTINATION", "Destination proposal did not pass QA.")
            else:
                neg.add("WAIT_FOR_DELIVERY", "Hold source negative until a fresh snapshot and delivery evidence confirm the new exact.", "DEFERRED")
            continue
        if norm(dest.get("State")) != "enabled" or not active_destination(dest):
            neg.add("HARVEST_INACTIVE", "Read-back destination is not enabled and eligible for static delivery checks.")
        if any(negative_scope(n, dest) and blocked_by(dest.get("Keyword Text"), n) for n in negatives + newneg if n is not neg and norm(n.get("State")) == "enabled"):
            neg.add("HARVEST_BLOCKED", "Destination is blocked by an existing/proposed negative.")
        evidence = link.get("delivery_evidence", {})
        try:
            checked_at = instant(evidence["checked_at"])
            age = (now - checked_at).total_seconds() / 3600
            start, end = instant(evidence["window_start"]), instant(evidence["window_end"])
            window_age = (now - start).total_seconds() / 3600
            impressions = amount(evidence["impressions"])
            valid = (str(evidence["keyword_id"]) == dest.get("Keyword ID") and
                     impressions == impressions.to_integral_value() and evidence["budget_available"] is True and
                     isinstance(evidence["reference"], str) and bool(evidence["reference"].strip()) and
                     isinstance(evidence["eligible_asins"], list) and
                     bool(sourceproducts) and sourceproducts <= (enabled_products[dest.get("Ad Group ID")] & {str(a).upper() for a in evidence["eligible_asins"]}) and
                     start < end <= checked_at <= now and
                     0 <= age <= float(amount(rules["delivery_max_age_hours"])) and
                     0 <= window_age <= float(amount(rules["delivery_max_age_hours"])))
        except (KeyError, ValueError, TypeError, AttributeError):
            valid = False
        if not valid:
            neg.add("DELIVERY_EVIDENCE", "Supply recent keyword-level impressions, budget availability, eligible ASINs and an evidence reference.", "NEEDS_CONTEXT")
    for neg in newneg:
        if neg.number in linked or norm(neg.get("Match Type")) != "negative exact":
            continue
        source = set().union(*(v for g, v in products.items() if groups.get(g) and groups[g].get("Campaign ID") == neg.get("Campaign ID"))) if not neg.get("Ad Group ID") else products[neg.get("Ad Group ID")]
        if any(norm(p.get("Match Type")) == "exact" and norm(p.get("State")) != "archived" and norm(p.get("Keyword Text")) == norm(neg.get("Keyword Text")) and source & products[p.get("Ad Group ID")] and not negative_scope(neg, p) for p in positives + creates):
            neg.add("HARVEST_LINK_REQUIRED", "This looks like traffic routing; connect it to a destination and delivery evidence.", "NEEDS_CONTEXT")
    severities = [i["severity"] for i in global_issues] + [r.status for r in proposals]
    status = "BLOCKED" if "BLOCKED" in severities else "NEEDS_CONTEXT" if "NEEDS_CONTEXT" in severities else "PHASE1_READY" if "DEFERRED" in severities else "READY"
    ready = [r for r in proposals if r.status == "READY"]
    if not proposals or (not ready and status == "READY"):
        status = "NO_CHANGES"
    return {"version": VERSION, "status": status, "reviewed_at": now.isoformat(),
            "context": {k: manifest.get(k) for k in ["account_id", "marketplace", "currency", "snapshot_exported_at", "snapshot_complete"]},
            "rules": rules, "global_issues": global_issues,
            "new_exact_counts": dict(counts), "rows": proposals,
            "limitations": ["Literal exact/phrase conflicts do not model all Amazon close variants.",
                            "Export context and delivery evidence are user-supplied, not live API verification.",
                            "Checked workbook still requires review and validation in the current Amazon console."]}


def csv_safe(value):
    s = str(value)
    return "'" + s if s.lstrip().startswith(("=", "+", "-", "@")) else s


def write_csv(path, headers, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows({k: csv_safe(r.get(k, "")) for k in headers} for r in rows)


def workbook(path, rows):
    try:
        import openpyxl
    except ImportError as exc:
        raise InputError("Install openpyxl to produce checked-upload.xlsx.") from exc
    wb = openpyxl.Workbook()
    wb.properties.creator = "Radu Bobirnac · Profit First PPC"
    wb.properties.lastModifiedBy = "Profit First PPC"
    wb.properties.title = "PPC Bulk File QA · checked upload"
    wb.properties.description = "Prepared with amazon-ppc-bulk-qa. " + ATTRIBUTION["referral_url"]
    ws = wb.active
    ws.title = "Sponsored Products Campaigns"
    values = [UPLOAD] + [[r.payload.get(h, "") for h in UPLOAD] for r in rows]
    for rr in values:
        ws.append(rr)
        for c in ws[ws.max_row]:
            c.data_type = "s"
            c.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    check = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        actual = [[c.value or "" for c in row] for row in check.active.iter_rows()]
        if actual != values or any(c.data_type == "f" for row in check.active.iter_rows() for c in row):
            path.unlink()
            raise InputError("Upload workbook round-trip failed; no checked upload produced.")
    finally:
        check.close()


def run(manifest_path, out):
    manifest_path, out = Path(manifest_path).resolve(), Path(out).resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if not isinstance(manifest, dict):
        raise InputError("Manifest must be a JSON object.")
    paths = {k: (manifest_path.parent / manifest[k]).resolve() for k in ["snapshot", "changes"]}
    if out.exists():
        raise InputError("Output directory exists; choose a new directory to preserve prior results.")
    hashes = {k: sha(p) for k, p in paths.items()}
    hashes["manifest"] = sha(manifest_path)
    snapshot, sh, sn = read_table(paths["snapshot"], manifest.get("snapshot_sheet"))
    proposals, ph, pn = read_table(paths["changes"], manifest.get("changes_sheet"))
    result = analyze(snapshot, proposals, manifest, sh)
    if hashes != {**{k: sha(p) for k, p in paths.items()}, "manifest": sha(manifest_path)}:
        raise InputError("An input changed while reading; export and run again.")
    out.mkdir(parents=True)
    serial = {**result, "attribution": dict(ATTRIBUTION), "rows": [{"row": r.number, "sheet": r.sheet, "status": r.status, "issues": r.issues, "before": r.before, "proposal": r.data, "upload": r.payload} for r in proposals],
              "inputs": {k: {"path": str(p), "sha256": hashes[k]} for k, p in paths.items()},
              "manifest_sha256": hashes["manifest"], "notes": sn + pn, "upload": None}
    ready = [r for r in proposals if r.status == "READY"]
    if result["status"] in {"READY", "PHASE1_READY"} and ready:
        try:
            workbook(out / "checked-upload.xlsx", ready)
            serial["upload"] = {"file": "checked-upload.xlsx", "rows": len(ready), "sha256": sha(out / "checked-upload.xlsx")}
        except InputError as exc:
            serial["status"] = "BLOCKED"
            serial["global_issues"].append(issue("EXPORT_FAILED", str(exc)))
    exceptions = [{"row": "", **i} for i in serial["global_issues"]]
    exceptions += [{"row": r.number, **i} for r in proposals for i in r.issues]
    write_csv(out / "exceptions.csv", ["row", "severity", "code", "message"], exceptions)
    changes = []
    for r in proposals:
        for k in ["Bid", "Daily Budget", "Keyword Text", "Match Type", "State"]:
            if k in r.payload:
                changes.append({"row": r.number, "status": r.status, "entity": r.get("Entity"), "operation": r.get("Operation"), "campaign_id": r.get("Campaign ID"), "ad_group_id": r.get("Ad Group ID"), "keyword_id": r.get("Keyword ID"), "campaign_name": next((x.get("Campaign Name") for x in snapshot if norm(x.get("Entity")) == "campaign" and x.get("Campaign ID") == r.get("Campaign ID")), ""), "field": k, "before": r.before.get(k, ""), "after": r.payload[k]})
    write_csv(out / "changes.csv", ["row", "status", "entity", "operation", "campaign_id", "ad_group_id", "keyword_id", "campaign_name", "field", "before", "after"], changes)
    write_csv(out / "deferred.csv", ["row", "reason", "proposal_json"], [{"row": r.number, "reason": " | ".join(i["message"] for i in r.issues if i["severity"] == "DEFERRED"), "proposal_json": json.dumps(r.data, ensure_ascii=False)} for r in proposals if r.status == "DEFERRED"])
    (out / "report.json").write_text(json.dumps(serial, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    audit = {"version": VERSION, "manifest": manifest, "input_hashes": hashes, "snapshot_headers": sh, "proposal_headers": ph, "notes": sn + pn, "snapshot_rows": [{"row": r.number, "sheet": r.sheet, "values": r.data} for r in snapshot], "proposal_rows": [{"row": r.number, "sheet": r.sheet, "values": r.data} for r in proposals]}
    (out / "audit.json").write_text(json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    lines = ["# PPC Bulk File QA", "", f"Status: **{serial['status']}**", "", f"Account: {manifest.get('account_id', 'unknown')} · {manifest.get('marketplace', 'unknown')} · {manifest.get('currency', 'unknown')}", "", f"Upload: {serial['upload']['rows']} checked rows in checked-upload.xlsx." if serial["upload"] else "No upload workbook produced.", "", "Review changes.csv and exceptions.csv before using any workbook. Deferred source negatives are excluded.", "", "## Findings", ""]
    lines += [f"- Row {e['row'] or 'global'} — {e['severity']} / {e['code']}: {e['message']}" for e in exceptions] or ["- No issues found by the implemented checks."]
    lines += ["", "## Limits", ""] + [f"- {x}" for x in serial["limitations"]]
    lines += ["", "---", "", CREDIT]
    (out / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return serial


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    try:
        result = run(args.manifest, args.out)
    except (InputError, ValueError, KeyError, OSError) as exc:
        print(f"Input error: {exc}", file=sys.stderr)
        return 2
    print(json.dumps({"status": result["status"], "upload": result["upload"], "report": str(args.out / "report.md")}))
    return 2 if result["status"] == "BLOCKED" else 3 if result["status"] == "NEEDS_CONTEXT" else 0


if __name__ == "__main__":
    sys.exit(main())
